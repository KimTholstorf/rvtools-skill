#!/usr/bin/env python3
"""Parse an RVTools workbook into a deterministic analysis payload."""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import sys
from collections import Counter
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

SCRIPT_DIR = Path(__file__).resolve().parent
if str(SCRIPT_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPT_DIR))

from rvtools.migration import TARGET_IDS, assess_migration


SCHEMA_VERSION = "2.0"
ANALYSIS_SHEETS = (
    "vHost",
    "vCluster",
    "vDatastore",
    "vDisk",
    "vNetwork",
    "vSnapshot",
    "vTools",
    "vUSB",
    "vCD",
    "dvPort",
    "vSC_VMK",
    "vHealth",
)


def _key(value):
    return re.sub(r"[^a-z0-9]+", "", str(value or "").casefold())


def _number(value, default=0):
    if value in (None, ""):
        return default
    try:
        return float(str(value).replace(",", "").strip())
    except (TypeError, ValueError):
        return default


def _whole_or_float(value):
    number = _number(value)
    return int(number) if float(number).is_integer() else number


def _truthy(value):
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value != 0
    return str(value or "").strip().casefold() in {"true", "yes", "y", "1", "enabled", "connected"}


def _optional_truthy(value):
    if value in (None, ""):
        return None
    return _truthy(value)


def _serializable(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return value


def _datetime_value(value):
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    text = str(value or "").strip()
    if not text:
        return None
    try:
        return datetime.fromisoformat(text)
    except ValueError:
        pass
    for pattern in ("%Y-%m-%d %H:%M:%S", "%m/%d/%Y %H:%M:%S", "%d/%m/%Y %H:%M:%S"):
        try:
            return datetime.strptime(text, pattern)
        except ValueError:
            continue
    return None


def _read_rows(workbook, sheet_name):
    if sheet_name not in workbook.sheetnames:
        return []
    worksheet = workbook[sheet_name]
    rows = worksheet.iter_rows(values_only=True)
    headers = next(rows, ())
    normalized = [_key(header) for header in headers]
    return [
        {normalized[index]: value for index, value in enumerate(row) if index < len(normalized) and normalized[index]}
        for row in rows
        if any(value not in (None, "") for value in row)
    ]


def _value(row, *headers, default=None):
    for header in headers:
        key = _key(header)
        if key in row and row[key] not in (None, ""):
            return row[key]
    return default


def _counter(rows, *headers):
    values = (_value(row, *headers) for row in rows)
    return dict(sorted(Counter(str(value).strip() for value in values if value not in (None, "")).items()))


def _sum(rows, *headers):
    return _whole_or_float(sum(_number(_value(row, *headers)) for row in rows))


def _metadata(rows):
    row = rows[0] if rows else {}
    return {
        "rvtools_version": _serializable(_value(row, "RVTools version")),
        "exported_at": _serializable(_value(row, "xlsx creation datetime")),
    }


def _vm_name(row):
    return str(_value(row, "VM", default="(unnamed VM)"))


def _workload_rows(rows):
    return [row for row in rows if not _truthy(_value(row, "Template"))]


def _guest_os(row):
    tools_os = str(_value(row, "OS according to the VMware Tools", default="")).strip()
    configured_os = str(
        _value(row, "OS according to the configuration file", default="")
    ).strip()
    if tools_os:
        return tools_os, "vmware_tools"
    if configured_os:
        return configured_os, "configuration"
    return "", "unknown"


def _guest_os_family(value):
    text = str(value or "").strip().casefold()
    if not text:
        return "unknown"
    if "windows" in text:
        return "windows"
    linux_markers = (
        "linux",
        "ubuntu",
        "debian",
        "centos",
        "red hat",
        "rhel",
        "suse",
        "photon",
        "rocky",
        "alma",
    )
    if any(marker in text for marker in linux_markers):
        return "linux"
    if "darwin" in text or "mac os" in text or "macos" in text:
        return "macos"
    if "freebsd" in text:
        return "freebsd"
    if "solaris" in text:
        return "solaris"
    return "other"


def _guest_os_facts(workloads):
    values = [_guest_os(row) for row in workloads]
    return {
        "families": dict(sorted(Counter(_guest_os_family(value) for value, _ in values).items())),
        "sources": dict(sorted(Counter(source for _, source in values).items())),
    }


def _cpu_vendor(value):
    text = str(value or "").strip()
    normalized = text.casefold()
    if not text:
        return "Unknown"
    if "intel" in normalized:
        return "Intel"
    if "amd" in normalized or "epyc" in normalized:
        return "AMD"
    if "ampere" in normalized:
        return "Ampere"
    if "ibm" in normalized or "powerpc" in normalized:
        return "IBM"
    return re.sub(r"[^A-Za-z0-9_-]+", "", text.split()[0]) or "Other"


def _host_hardware_facts(hosts):
    models = Counter()
    available = Counter()
    active = Counter()
    for row in hosts:
        vendor = str(_value(row, "Vendor", default="")).strip()
        model = str(_value(row, "Model", default="")).strip()
        if vendor or model:
            models[(vendor, model)] += 1
        for counter, header in ((available, "HT Available"), (active, "HT Active")):
            status = _optional_truthy(_value(row, header))
            counter["unknown" if status is None else str(status).casefold()] += 1
    return {
        "models": [
            {"vendor": vendor, "model": model, "hosts": count}
            for (vendor, model), count in sorted(
                models.items(), key=lambda item: (item[0][0].casefold(), item[0][1].casefold())
            )
        ],
        "hyperthreading": {
            "available": {key: available[key] for key in ("true", "false", "unknown")},
            "active": {key: active[key] for key in ("true", "false", "unknown")},
        },
    }


def _ratio(numerator, denominator):
    return round(numerator / denominator, 2) if denominator > 0 else None


def _overcommit_summary(rows):
    workloads = _workload_rows(rows.get("vInfo", []))
    hosts = rows.get("vHost", [])

    def cluster_name(row):
        return str(_value(row, "Cluster", default="")).strip() or "(unassigned)"

    clusters = sorted({cluster_name(row) for row in workloads + hosts})
    summaries = []
    warnings = []
    for sheet_name, source_rows in (("vInfo", workloads), ("vHost", hosts)):
        for source_header in ("VI SDK UUID", "VI SDK Server"):
            sources = {
                str(_value(row, source_header)).strip()
                for row in source_rows
                if _value(row, source_header) not in (None, "")
            }
            if len(sources) > 1:
                warnings.append(
                    {
                        "code": "multi_vcenter_sources",
                        "sheet": sheet_name,
                        "source_count": len(sources),
                        "message": "Cluster names may overlap across vCenters; overcommit ratios are aggregated by cluster name.",
                    }
                )
            if sources:
                break
    for cluster in clusters:
        cluster_hosts = [row for row in hosts if cluster_name(row) == cluster]
        powered_on = [
            row
            for row in workloads
            if cluster_name(row) == cluster
            and str(_value(row, "Powerstate", default="")).strip().casefold() == "poweredon"
        ]
        powered_off = [
            row
            for row in workloads
            if cluster_name(row) == cluster
            and str(_value(row, "Powerstate", default="")).strip().casefold() == "poweredoff"
        ]
        configured_vcpus = _sum(powered_on, "CPUs")
        physical_cores = _sum(cluster_hosts, "# Cores", "Cores")
        configured_memory_mib = _sum(powered_on, "Memory", "Size MiB")
        physical_memory_mib = _sum(cluster_hosts, "# Memory", "Memory")
        if powered_on and physical_cores <= 0:
            warnings.append(
                {
                    "code": "missing_physical_cpu_capacity",
                    "cluster": cluster,
                    "message": "CPU overcommit ratio is unavailable because physical core capacity is missing or zero.",
                }
            )
        if powered_on and physical_memory_mib <= 0:
            warnings.append(
                {
                    "code": "missing_physical_memory_capacity",
                    "cluster": cluster,
                    "message": "Memory overcommit ratio is unavailable because physical memory capacity is missing or zero.",
                }
            )
        summaries.append(
            {
                "cluster": cluster,
                "host_count": len(cluster_hosts),
                "powered_on_vms": len(powered_on),
                "configured_vcpus": configured_vcpus,
                "physical_cores": physical_cores,
                "cpu_ratio": _ratio(configured_vcpus, physical_cores),
                "configured_memory_mib": configured_memory_mib,
                "physical_memory_mib": physical_memory_mib,
                "memory_ratio": _ratio(configured_memory_mib, physical_memory_mib),
                "powered_off_vms": len(powered_off),
                "powered_off_vcpus": _sum(powered_off, "CPUs"),
                "powered_off_memory_mib": _sum(powered_off, "Memory", "Size MiB"),
            }
        )

    def total(field):
        return _whole_or_float(sum(_number(item[field]) for item in summaries))

    overall = {
        "host_count": total("host_count"),
        "powered_on_vms": total("powered_on_vms"),
        "configured_vcpus": total("configured_vcpus"),
        "physical_cores": total("physical_cores"),
        "configured_memory_mib": total("configured_memory_mib"),
        "physical_memory_mib": total("physical_memory_mib"),
        "powered_off_vms": total("powered_off_vms"),
        "powered_off_vcpus": total("powered_off_vcpus"),
        "powered_off_memory_mib": total("powered_off_memory_mib"),
    }
    cpu_coverage_complete = all(
        item["powered_on_vms"] == 0 or item["physical_cores"] > 0 for item in summaries
    )
    memory_coverage_complete = all(
        item["powered_on_vms"] == 0 or item["physical_memory_mib"] > 0 for item in summaries
    )
    overall["cpu_ratio"] = (
        _ratio(overall["configured_vcpus"], overall["physical_cores"])
        if cpu_coverage_complete
        else None
    )
    overall["memory_ratio"] = (
        _ratio(overall["configured_memory_mib"], overall["physical_memory_mib"])
        if memory_coverage_complete
        else None
    )
    return {
        "methodology": {
            "vm_scope": "powered_on_workloads",
            "templates": "excluded",
            "cpu_ratio": "configured_vcpus / physical_cores",
            "memory_ratio": "configured_memory_mib / physical_memory_mib",
            "powered_off_vms": "reported_separately",
        },
        "clusters": summaries,
        "overall": overall,
        "warnings": warnings,
    }


def _hardware_version(value):
    match = re.search(r"(\d+)", str(value or ""))
    return int(match.group(1)) if match else None


def _add_detection(detections, *, detection_id, category, severity, title, matches, max_examples):
    if not matches:
        return
    detections.append(
        {
            "id": detection_id,
            "category": category,
            "severity": severity,
            "title": title,
            "tags": [f"source.{category}"],
            "count": len(matches),
            "examples": matches[:max_examples],
            "examples_truncated": len(matches) > max_examples,
        }
    )


def _vm_detections(rows, max_examples):
    detections = []
    workloads = [row for row in rows["vInfo"] if not _truthy(_value(row, "Template"))]
    def vm_matches(predicate, detail=None):
        matches = []
        for row in workloads:
            if predicate(row):
                item = {"vm": _vm_name(row)}
                if detail:
                    item.update(detail(row))
                matches.append(item)
        return matches

    _add_detection(
        detections,
        detection_id="vm_suspended",
        category="power_state",
        severity="high",
        title="Suspended virtual machines",
        matches=vm_matches(lambda row: "suspend" in str(_value(row, "Powerstate", default="")).casefold()),
        max_examples=max_examples,
    )
    _add_detection(
        detections,
        detection_id="consolidation_needed",
        category="storage",
        severity="medium",
        title="Virtual machines require disk consolidation",
        matches=vm_matches(lambda row: _truthy(_value(row, "Consolidation Needed"))),
        max_examples=max_examples,
    )
    _add_detection(
        detections,
        detection_id="vm_cpu_large",
        category="sizing",
        severity="high",
        title="Virtual machines exceed the default 128-vCPU review threshold",
        matches=vm_matches(
            lambda row: _number(_value(row, "CPUs")) > 128,
            lambda row: {"vcpus": _whole_or_float(_value(row, "CPUs"))},
        ),
        max_examples=max_examples,
    )
    _add_detection(
        detections,
        detection_id="vm_memory_large",
        category="sizing",
        severity="high",
        title="Virtual machines exceed the default 1-TiB memory review threshold",
        matches=vm_matches(
            lambda row: _number(_value(row, "Memory", "Size MiB")) > 1048576,
            lambda row: {"memory_mib": _whole_or_float(_value(row, "Memory", "Size MiB"))},
        ),
        max_examples=max_examples,
    )
    _add_detection(
        detections,
        detection_id="vm_provisioned_storage_large",
        category="sizing",
        severity="medium",
        title="Virtual machines exceed the 10-TiB provisioned-storage review threshold",
        matches=vm_matches(
            lambda row: _number(_value(row, "Provisioned MiB")) > 10 * 1024 * 1024,
            lambda row: {
                "provisioned_mib": _whole_or_float(_value(row, "Provisioned MiB")),
                "in_use_mib": _whole_or_float(_value(row, "In Use MiB")),
            },
        ),
        max_examples=max_examples,
    )
    _add_detection(
        detections,
        detection_id="legacy_vm_hardware",
        category="compatibility",
        severity="medium",
        title="Virtual machines use hardware version older than vmx-14",
        matches=vm_matches(
            lambda row: (_hardware_version(_value(row, "HW version")) or 999) < 14,
            lambda row: {"hardware_version": str(_value(row, "HW version"))},
        ),
        max_examples=max_examples,
    )
    _add_detection(
        detections,
        detection_id="fault_tolerance_enabled",
        category="migration_compatibility",
        severity="high",
        title="Fault Tolerance is configured",
        matches=vm_matches(
            lambda row: str(_value(row, "FT State", default="")).strip().casefold()
            not in {"", "notconfigured", "not configured", "disabled", "off"}
        ),
        max_examples=max_examples,
    )
    oracle_pattern = re.compile(r"\boracle\b", re.IGNORECASE)
    _add_detection(
        detections,
        detection_id="oracle_workload",
        category="licensing",
        severity="medium",
        title="Potential Oracle workloads require licensing review",
        matches=vm_matches(
            lambda row: any(
                oracle_pattern.search(str(_value(row, header, default="")))
                for header in (
                    "VM",
                    "Annotation",
                    "OS according to the configuration file",
                    "OS according to the VMware Tools",
                )
            )
        ),
        max_examples=max_examples,
    )
    secret_pattern = re.compile(r"\b(password|passwd|pwd|secret|token|api[ _-]?key)\b\s*[:=]", re.IGNORECASE)
    secret_matches = []
    for row in workloads:
        if secret_pattern.search(str(_value(row, "Annotation", default=""))):
            secret_matches.append({"vm": _vm_name(row), "field": "vInfo.Annotation"})
    for row in _workload_rows(rows.get("vSnapshot", [])):
        for header in ("Name", "Description", "Annotation"):
            if secret_pattern.search(str(_value(row, header, default=""))):
                secret_matches.append({"vm": _vm_name(row), "field": f"vSnapshot.{header}"})
    _add_detection(
        detections,
        detection_id="possible_cleartext_secret",
        category="security",
        severity="high",
        title="Possible clear-text credentials or tokens",
        matches=secret_matches,
        max_examples=max_examples,
    )
    _add_detection(
        detections,
        detection_id="tools_not_running",
        category="guest_operations",
        severity="medium",
        title="VMware Tools is not running or not installed",
        matches=[
            {"vm": _vm_name(row), "status": str(_value(row, "Tools"))}
            for row in _workload_rows(rows.get("vTools", []))
            if str(_value(row, "Powerstate", default="")).strip().casefold() == "poweredon"
            and any(token in str(_value(row, "Tools", default="")).casefold() for token in ("notrunning", "not running", "notinstalled", "not installed"))
        ],
        max_examples=max_examples,
    )
    _add_detection(
        detections,
        detection_id="connected_usb",
        category="device",
        severity="high",
        title="Connected USB devices",
        matches=[
            {"vm": _vm_name(row)}
            for row in _workload_rows(rows.get("vUSB", []))
            if _truthy(_value(row, "Connected"))
        ],
        max_examples=max_examples,
    )
    _add_detection(
        detections,
        detection_id="connected_cdrom",
        category="device",
        severity="medium",
        title="Connected CD/DVD devices",
        matches=[
            {"vm": _vm_name(row)}
            for row in _workload_rows(rows.get("vCD", []))
            if _truthy(_value(row, "Connected"))
        ],
        max_examples=max_examples,
    )
    _add_detection(
        detections,
        detection_id="cdrom_starts_connected",
        category="device",
        severity="low",
        title="Disconnected CD/DVD devices are configured to connect at power-on",
        matches=[
            {"vm": _vm_name(row)}
            for row in _workload_rows(rows.get("vCD", []))
            if not _truthy(_value(row, "Connected"))
            and _truthy(_value(row, "Starts Connected"))
        ],
        max_examples=max_examples,
    )
    return detections


def _storage_detections(rows, max_examples, reference_time):
    detections = []
    disk_rows = _workload_rows(rows.get("vDisk", []))

    stale_snapshots = []
    for row in _workload_rows(rows.get("vSnapshot", [])):
        created = _datetime_value(_value(row, "Date / time"))
        if created is None:
            continue
        age_days = max(0, (reference_time - created).days)
        if age_days >= 3:
            stale_snapshots.append(
                {
                    "vm": _vm_name(row),
                    "created_at": created.isoformat(),
                    "age_days": age_days,
                    "size_mib": _whole_or_float(_value(row, "Size MiB (total)")),
                }
            )
    _add_detection(
        detections,
        detection_id="stale_snapshot",
        category="snapshot",
        severity="high",
        title="Snapshots are at least 72 hours old",
        matches=stale_snapshots,
        max_examples=max_examples,
    )

    raw_disks = []
    for row in disk_rows:
        compatibility = str(_value(row, "Raw Comp. Mode", default="")).strip()
        if _truthy(_value(row, "Raw")) or compatibility:
            raw_disks.append(
                {
                    "vm": _vm_name(row),
                    "disk": str(_value(row, "Disk", default="(unnamed disk)")),
                    "compatibility_mode": compatibility or "unspecified",
                }
            )
    _add_detection(
        detections,
        detection_id="raw_device_mapping",
        category="storage_compatibility",
        severity="high",
        title="Raw Device Mapping disks",
        matches=raw_disks,
        max_examples=max_examples,
    )
    _add_detection(
        detections,
        detection_id="independent_disk",
        category="storage_compatibility",
        severity="high",
        title="Independent-mode virtual disks",
        matches=[
            {
                "vm": _vm_name(row),
                "disk": str(_value(row, "Disk", default="(unnamed disk)")),
                "mode": str(_value(row, "Disk Mode")),
            }
            for row in disk_rows
            if "independent" in str(_value(row, "Disk Mode", default="")).casefold()
        ],
        max_examples=max_examples,
    )
    no_sharing = {"", "nosharing", "no sharing", "sharingnone", "none", "false"}
    _add_detection(
        detections,
        detection_id="shared_disk",
        category="storage_compatibility",
        severity="high",
        title="Shared or multi-writer virtual disks",
        matches=[
            {
                "vm": _vm_name(row),
                "disk": str(_value(row, "Disk", default="(unnamed disk)")),
                "sharing_mode": str(_value(row, "Sharing mode", default="")),
            }
            for row in disk_rows
            if str(_value(row, "Sharing mode", default="")).strip().casefold() not in no_sharing
            or str(_value(row, "Shared Bus", default="")).strip().casefold() not in no_sharing
        ],
        max_examples=max_examples,
    )

    low_free = []
    overprovisioned = []
    for row in rows.get("vDatastore", []):
        name = str(_value(row, "Name", default="(unnamed datastore)"))
        free_percent = _number(_value(row, "Free %"), default=-1)
        if 0 <= free_percent <= 1:
            free_percent *= 100
        if 0 <= free_percent < 10:
            low_free.append({"datastore": name, "free_percent": round(free_percent, 2)})
        capacity = _number(_value(row, "Capacity MiB"))
        provisioned = _number(_value(row, "Provisioned MiB"))
        if capacity > 0 and provisioned > capacity:
            overprovisioned.append(
                {
                    "datastore": name,
                    "provisioned_percent": round((provisioned / capacity) * 100, 1),
                }
            )
    _add_detection(
        detections,
        detection_id="datastore_low_free_space",
        category="capacity",
        severity="high",
        title="Datastores have less than 10% free space",
        matches=low_free,
        max_examples=max_examples,
    )
    _add_detection(
        detections,
        detection_id="datastore_overprovisioned",
        category="capacity",
        severity="medium",
        title="Datastore provisioned capacity exceeds physical capacity",
        matches=overprovisioned,
        max_examples=max_examples,
    )
    return detections


def _network_detections(rows, max_examples):
    detections = []
    network_rows = _workload_rows(rows.get("vNetwork", []))
    vmkernel_rows = rows.get("vSC_VMK", [])
    vmkernel_port_groups = {
        str(_value(row, "Port Group", default="")).strip().casefold()
        for row in vmkernel_rows
        if str(_value(row, "Port Group", default="")).strip()
    }
    vmkernel_host_port_groups = {
        (
            str(_value(row, "Host", default="")).strip().casefold(),
            str(_value(row, "Port Group", default="")).strip().casefold(),
        )
        for row in vmkernel_rows
        if str(_value(row, "Host", default="")).strip()
        and str(_value(row, "Port Group", default="")).strip()
    }
    unscoped_vmk_port_groups = {
        str(_value(row, "Port Group", default="")).strip().casefold()
        for row in vmkernel_rows
        if not str(_value(row, "Host", default="")).strip()
        and str(_value(row, "Port Group", default="")).strip()
    }
    standard_switch = []
    vmkernel_attachment = []
    for row in network_rows:
        switch = str(_value(row, "Switch", default="")).strip()
        network = str(_value(row, "Network", default="")).strip()
        host = str(_value(row, "Host", default="")).strip()
        example = {"vm": _vm_name(row), "network": network, "switch": switch}
        if switch.casefold().startswith("vswitch"):
            standard_switch.append(example)
        is_vmk_network = (
            (host.casefold(), network.casefold()) in vmkernel_host_port_groups
            or network.casefold() in unscoped_vmk_port_groups
            or (not host and network.casefold() in vmkernel_port_groups)
        )
        if is_vmk_network:
            vmkernel_attachment.append(example)
    _add_detection(
        detections,
        detection_id="standard_vswitch_attachment",
        category="network_compatibility",
        severity="medium",
        title="VM network adapters use standard vSwitches",
        matches=standard_switch,
        max_examples=max_examples,
    )
    _add_detection(
        detections,
        detection_id="vmkernel_network_attachment",
        category="network_isolation",
        severity="high",
        title="VM network adapters use VMkernel port groups",
        matches=vmkernel_attachment,
        max_examples=max_examples,
    )

    dvports = rows.get("dvPort", [])

    def port_example(row):
        return {
            "port_group": str(_value(row, "Port", default="(unnamed port group)")),
            "switch": str(_value(row, "Switch", default="")),
        }

    def vlan_zero_or_empty(row):
        vlan = _value(row, "VLAN", default=None)
        return vlan in (None, "") or _number(vlan, default=-1) == 0

    _add_detection(
        detections,
        detection_id="dvport_vlan_zero",
        category="network_security",
        severity="medium",
        title="Distributed port groups use VLAN 0 or have no VLAN ID",
        matches=[port_example(row) for row in dvports if vlan_zero_or_empty(row)],
        max_examples=max_examples,
    )
    _add_detection(
        detections,
        detection_id="dvport_promiscuous_mode",
        category="network_security",
        severity="high",
        title="Distributed port groups allow promiscuous mode",
        matches=[port_example(row) for row in dvports if _truthy(_value(row, "Allow Promiscuous"))],
        max_examples=max_examples,
    )
    _add_detection(
        detections,
        detection_id="dvport_mac_changes",
        category="network_security",
        severity="medium",
        title="Distributed port groups allow MAC address changes",
        matches=[port_example(row) for row in dvports if _truthy(_value(row, "Mac Changes"))],
        max_examples=max_examples,
    )
    _add_detection(
        detections,
        detection_id="dvport_forged_transmits",
        category="network_security",
        severity="medium",
        title="Distributed port groups allow forged transmits",
        matches=[port_example(row) for row in dvports if _truthy(_value(row, "Forged Transmits"))],
        max_examples=max_examples,
    )
    _add_detection(
        detections,
        detection_id="dvport_ephemeral_binding",
        category="network_compatibility",
        severity="medium",
        title="Distributed port groups use ephemeral binding",
        matches=[
            port_example(row)
            for row in dvports
            if "ephemeral" in str(_value(row, "Type", default="")).casefold()
        ],
        max_examples=max_examples,
    )
    return detections


def _environment_detections(rows, max_examples):
    detections = []
    hosts = rows.get("vHost", [])
    _add_detection(
        detections,
        detection_id="non_intel_host",
        category="cpu_compatibility",
        severity="medium",
        title="Hosts use non-Intel processors",
        matches=[
            {
                "host": str(_value(row, "Host", default="(unnamed host)")),
                "cpu_model": str(_value(row, "CPU Model")),
            }
            for row in hosts
            if str(_value(row, "CPU Model", default="")).strip()
            and "intel" not in str(_value(row, "CPU Model", default="")).casefold()
        ],
        max_examples=max_examples,
    )
    version_counts = _counter(hosts, "ESX Version")
    _add_detection(
        detections,
        detection_id="esxi_version_spread",
        category="lifecycle",
        severity="medium",
        title="Multiple ESXi versions are present",
        matches=[{"version": version, "hosts": count} for version, count in version_counts.items()]
        if len(version_counts) > 1
        else [],
        max_examples=max_examples,
    )
    _add_detection(
        detections,
        detection_id="host_cpu_pressure",
        category="capacity",
        severity="medium",
        title="Hosts exceed 80% CPU utilization",
        matches=[
            {
                "host": str(_value(row, "Host", default="(unnamed host)")),
                "cpu_usage_percent": _whole_or_float(_value(row, "CPU usage %")),
            }
            for row in hosts
            if _number(_value(row, "CPU usage %")) > 80
        ],
        max_examples=max_examples,
    )
    _add_detection(
        detections,
        detection_id="host_memory_pressure",
        category="capacity",
        severity="high",
        title="Hosts exceed 90% memory utilization",
        matches=[
            {
                "host": str(_value(row, "Host", default="(unnamed host)")),
                "memory_usage_percent": _whole_or_float(_value(row, "Memory usage %")),
            }
            for row in hosts
            if _number(_value(row, "Memory usage %")) > 90
        ],
        max_examples=max_examples,
    )
    _add_detection(
        detections,
        detection_id="host_hyperthreading_inactive",
        category="hardware_configuration",
        severity="medium",
        title="Hosts support Hyper-Threading but do not have it active",
        matches=[
            {
                "host": str(_value(row, "Host", default="(unnamed host)")),
                "vendor": str(_value(row, "Vendor", default="")),
                "model": str(_value(row, "Model", default="")),
                "ht_available": True,
                "ht_active": False,
            }
            for row in hosts
            if _optional_truthy(_value(row, "HT Available")) is True
            and _optional_truthy(_value(row, "HT Active")) is False
        ],
        max_examples=max_examples,
    )
    error_types = {"error", "critical", "red", "alert"}
    _add_detection(
        detections,
        detection_id="rvtools_health_error",
        category="health",
        severity="high",
        title="RVTools health checks report errors",
        matches=[
            {
                "check": str(_value(row, "Name", default="(unnamed health check)")),
                "message_type": str(_value(row, "Message type")),
            }
            for row in rows.get("vHealth", [])
            if str(_value(row, "Message type", default="")).strip().casefold() in error_types
        ],
        max_examples=max_examples,
    )
    return detections


def analyze_workbook(
    path,
    *,
    max_examples=25,
    target=None,
    target_node=None,
    target_region=None,
    target_cpu_vendor=None,
):
    """Return the normalized analysis payload for an RVTools workbook."""
    path = Path(path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheetnames = workbook.sheetnames
        if "vInfo" not in sheetnames:
            raise ValueError("Not an RVTools export: required sheet vInfo is missing")
        rows = {name: _read_rows(workbook, name) for name in sheetnames}
    finally:
        workbook.close()
    vm_rows = rows["vInfo"]
    workloads = [row for row in vm_rows if not _truthy(_value(row, "Template"))]
    templates = [row for row in vm_rows if _truthy(_value(row, "Template"))]
    hosts = rows.get("vHost", [])
    datastores = rows.get("vDatastore", [])
    disk_rows = rows.get("vDisk", [])
    workload_disks = _workload_rows(disk_rows)
    template_disks = [row for row in disk_rows if _truthy(_value(row, "Template"))]
    network_rows = rows.get("vNetwork", [])
    workload_networks = _workload_rows(network_rows)
    template_networks = [row for row in network_rows if _truthy(_value(row, "Template"))]
    snapshot_rows = rows.get("vSnapshot", [])
    workload_snapshots = _workload_rows(snapshot_rows)
    template_snapshots = [row for row in snapshot_rows if _truthy(_value(row, "Template"))]
    metadata = _metadata(rows.get("vMetaData", []))
    guest_os_facts = _guest_os_facts(workloads)
    host_hardware_facts = _host_hardware_facts(hosts)
    metadata_row = rows.get("vMetaData", [{}])[0] if rows.get("vMetaData") else {}
    reference_time = _datetime_value(_value(metadata_row, "xlsx creation datetime")) or datetime.now()
    warnings = [
        {
            "code": "missing_sheet",
            "sheet": sheet,
            "message": f"{sheet} is absent; related facts and detections are incomplete.",
        }
        for sheet in ANALYSIS_SHEETS
        if sheet not in sheetnames
    ]

    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    result = {
        "schema_version": SCHEMA_VERSION,
        "source": {
            "file": path.name,
            "sha256": digest,
            "rvtools_version": metadata["rvtools_version"],
            "exported_at": metadata["exported_at"],
            "sheets_present": sheetnames,
        },
        "inventory": {
            "vms": len(workloads),
            "templates": len(templates),
            "hosts": len(hosts),
            "clusters": len(rows.get("vCluster", [])),
            "datastores": len(datastores),
            "virtual_disks": len(workload_disks),
            "template_virtual_disks": len(template_disks),
            "virtual_disks_total": len(disk_rows),
            "network_adapters": len(workload_networks),
            "template_network_adapters": len(template_networks),
            "network_adapters_total": len(network_rows),
            "snapshots": len(workload_snapshots),
            "template_snapshots": len(template_snapshots),
            "snapshots_total": len(snapshot_rows),
        },
        "capacity_mib": {
            "vm_provisioned": _sum(workloads, "Provisioned MiB"),
            "vm_in_use": _sum(workloads, "In Use MiB"),
            "datastore_capacity": _sum(datastores, "Capacity MiB"),
            "datastore_provisioned": _sum(datastores, "Provisioned MiB"),
            "datastore_in_use": _sum(datastores, "In Use MiB"),
            "datastore_free": _sum(datastores, "Free MiB"),
        },
        "overcommit": _overcommit_summary(rows),
        "facts": {
            "power_states": _counter(workloads, "Powerstate"),
            "esxi_versions": _counter(hosts, "ESX Version"),
            "cpu_models": _counter(hosts, "CPU Model"),
            "cpu_vendors": dict(
                sorted(Counter(_cpu_vendor(_value(row, "CPU Model")) for row in hosts).items())
            ),
            "host_hardware_models": host_hardware_facts["models"],
            "host_hyperthreading": host_hardware_facts["hyperthreading"],
            "vm_hardware_versions": _counter(workloads, "HW version"),
            "guest_os_families": guest_os_facts["families"],
            "guest_os_sources": guest_os_facts["sources"],
            "vm_counts_by_cluster": _counter(workloads, "Cluster"),
        },
        "detections": _vm_detections(rows, max_examples)
        + _storage_detections(rows, max_examples, reference_time)
        + _network_detections(rows, max_examples)
        + _environment_detections(rows, max_examples),
        "warnings": warnings,
    }
    if target:
        result["migration"] = assess_migration(
            rows,
            target,
            target_node=target_node,
            target_region=target_region,
            target_cpu_vendor=target_cpu_vendor,
        )
    return result


def _positive_integer(value):
    number = int(value)
    if number < 1:
        raise argparse.ArgumentTypeError("must be at least 1")
    return number


def main(argv=None):
    parser = argparse.ArgumentParser(
        description="Parse an RVTools .xlsx export into deterministic JSON facts and detections."
    )
    parser.add_argument("workbook", type=Path, help="Path to the RVTools .xlsx export")
    parser.add_argument("-o", "--output", type=Path, help="Write JSON to this file instead of stdout")
    parser.add_argument("--pretty", action="store_true", help="Indent the JSON output")
    parser.add_argument(
        "--max-examples",
        type=_positive_integer,
        default=25,
        help="Maximum affected-object examples retained per detection (default: 25)",
    )
    parser.add_argument(
        "--target",
        choices=TARGET_IDS,
        help="Add an HCX migration assessment for this cloud VMware target",
    )
    parser.add_argument("--target-node", help="Select a provider node or host type")
    parser.add_argument("--target-region", help="Record the intended cloud region")
    parser.add_argument(
        "--target-cpu-vendor",
        choices=("Intel", "AMD"),
        help="Select the OCVS target CPU vendor for live-migration screening",
    )
    args = parser.parse_args(argv)
    if args.output and args.output.resolve() == args.workbook.resolve():
        parser.exit(2, "error: output path must not overwrite the source workbook\n")
    try:
        payload = analyze_workbook(
            args.workbook,
            max_examples=args.max_examples,
            target=args.target,
            target_node=args.target_node,
            target_region=args.target_region,
            target_cpu_vendor=args.target_cpu_vendor,
        )
    except (FileNotFoundError, OSError, ValueError) as exc:
        parser.exit(2, f"error: {exc}\n")

    serialized = json.dumps(payload, indent=2 if args.pretty else None, sort_keys=args.pretty)
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        descriptor = os.open(args.output, os.O_WRONLY | os.O_CREAT | os.O_TRUNC, 0o600)
        if hasattr(os, "fchmod"):
            os.fchmod(descriptor, 0o600)
        with os.fdopen(descriptor, "w", encoding="utf-8") as output_file:
            output_file.write(serialized + "\n")
    else:
        sys.stdout.write(serialized + "\n")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
