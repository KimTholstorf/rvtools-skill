#!/usr/bin/env python3
"""Safely query a local RVTools workbook through an allowlisted inventory index."""

from __future__ import annotations

import argparse
import hashlib
import importlib.util
import json
import math
import os
import sqlite3
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook


PARSER_PATH = Path(__file__).with_name("parse_rvtools.py")
PARSER_SPEC = importlib.util.spec_from_file_location("rvtools_query_parser", PARSER_PATH)
RVTOOLS = importlib.util.module_from_spec(PARSER_SPEC)
assert PARSER_SPEC.loader is not None
PARSER_SPEC.loader.exec_module(RVTOOLS)

from rvtools.licensing import host_core_requirement

INDEX_SCHEMA_VERSION = "8"
INDEXED_SHEETS = (
    "vInfo",
    "vHost",
    "vTools",
    "vDatastore",
    "vDisk",
    "vNetwork",
    "vSnapshot",
    "vCD",
    "vUSB",
    "dvPort",
    "vLicense",
)

ENTITY_SCHEMAS = {
    "vm": {
        "vm": "TEXT",
        "power_state": "TEXT",
        "template": "INTEGER",
        "cluster": "TEXT",
        "host": "TEXT",
        "guest_os": "TEXT",
        "guest_os_family": "TEXT",
        "guest_os_source": "TEXT",
        "cpus": "REAL",
        "memory_mib": "REAL",
        "provisioned_mib": "REAL",
        "in_use_mib": "REAL",
        "hardware_version": "TEXT",
        "tools_status": "TEXT",
        "consolidation_needed": "INTEGER",
    },
    "host": {
        "host": "TEXT",
        "cluster": "TEXT",
        "vendor": "TEXT",
        "model": "TEXT",
        "cpu_vendor": "TEXT",
        "cpu_model": "TEXT",
        "ht_available": "INTEGER",
        "ht_active": "INTEGER",
        "cpu_sockets": "REAL",
        "cores_per_cpu": "REAL",
        "esxi_version": "TEXT",
        "cores": "REAL",
        "cpu_speed_mhz": "REAL",
        "memory_mib": "REAL",
        "cpu_usage_percent": "REAL",
        "memory_usage_percent": "REAL",
        "bios_vendor": "TEXT",
        "bios_version": "TEXT",
        "bios_date": "TEXT",
    },
    "cluster": {
        "cluster": "TEXT",
        "host_count": "INTEGER",
        "powered_on_vms": "INTEGER",
        "configured_vcpus": "REAL",
        "physical_cores": "REAL",
        "cpu_ratio": "REAL",
        "configured_memory_mib": "REAL",
        "physical_memory_mib": "REAL",
        "memory_ratio": "REAL",
        "powered_off_vms": "INTEGER",
        "powered_off_vcpus": "REAL",
        "powered_off_memory_mib": "REAL",
    },
    "datastore": {
        "datastore": "TEXT",
        "type": "TEXT",
        "cluster": "TEXT",
        "capacity_mib": "REAL",
        "provisioned_mib": "REAL",
        "in_use_mib": "REAL",
        "free_mib": "REAL",
        "free_percent": "REAL",
        "accessible": "INTEGER",
    },
    "license": {
        "name": "TEXT",
        "cost_unit": "TEXT",
        "total": "REAL",
        "used": "REAL",
        "expiration_date": "TEXT",
        "vi_sdk_server": "TEXT",
    },
    "vcf_license": {
        "host": "TEXT",
        "cluster": "TEXT",
        "cpu_sockets": "REAL",
        "cores_per_cpu": "REAL",
        "physical_cores": "REAL",
        "vcf_licensable_cores": "REAL",
        "core_minimum_adjustment": "REAL",
    },
    "vcf_license_summary": {
        "host_count": "INTEGER",
        "physical_cores": "REAL",
        "vcf_licensable_cores": "REAL",
        "core_calculation_complete": "INTEGER",
        "vsan_entitlement_tib": "REAL",
        "vsan_capacity_tib": "REAL",
        "vsan_capacity_evidence": "TEXT",
        "vsan_capacity_is_raw": "INTEGER",
        "vsan_addon_required_tib": "REAL",
        "vsan_entitlement_surplus_tib": "REAL",
    },
    "disk": {
        "vm": "TEXT",
        "template": "INTEGER",
        "power_state": "TEXT",
        "cluster": "TEXT",
        "host": "TEXT",
        "disk": "TEXT",
        "capacity_mib": "REAL",
        "raw": "INTEGER",
        "disk_mode": "TEXT",
        "sharing_mode": "TEXT",
        "raw_compatibility_mode": "TEXT",
    },
    "network": {
        "vm": "TEXT",
        "template": "INTEGER",
        "power_state": "TEXT",
        "network": "TEXT",
        "switch": "TEXT",
        "connected": "INTEGER",
        "cluster": "TEXT",
        "host": "TEXT",
    },
    "snapshot": {
        "vm": "TEXT",
        "template": "INTEGER",
        "cluster": "TEXT",
        "host": "TEXT",
        "created_at": "TEXT",
        "size_mib": "REAL",
        "power_state": "TEXT",
    },
    "dvport": {
        "port_group": "TEXT",
        "switch": "TEXT",
        "vlan": "TEXT",
        "allow_promiscuous": "INTEGER",
        "mac_changes": "INTEGER",
        "forged_transmits": "INTEGER",
        "binding_type": "TEXT",
    },
    "cdrom": {
        "vm": "TEXT",
        "template": "INTEGER",
        "power_state": "TEXT",
        "cluster": "TEXT",
        "host": "TEXT",
        "connected": "INTEGER",
        "starts_connected": "INTEGER",
        "device_type": "TEXT",
    },
    "usb": {
        "vm": "TEXT",
        "template": "INTEGER",
        "power_state": "TEXT",
        "cluster": "TEXT",
        "host": "TEXT",
        "connected": "INTEGER",
        "device_type": "TEXT",
    },
    "migration_method": {
        "target": "TEXT",
        "vm": "TEXT",
        "cluster": "TEXT",
        "host": "TEXT",
        "vcenter": "TEXT",
        "method": "TEXT",
        "status": "TEXT",
        "reason_ids": "TEXT",
    },
    "migration_finding": {
        "target": "TEXT",
        "finding_id": "TEXT",
        "category": "TEXT",
        "vm": "TEXT",
        "cluster": "TEXT",
        "host": "TEXT",
        "vcenter": "TEXT",
        "status": "TEXT",
        "methods": "TEXT",
        "summary": "TEXT",
    },
    "target_node": {
        "target": "TEXT",
        "node_type": "TEXT",
        "shape_series": "TEXT",
        "physical_cores": "REAL",
        "configured_physical_cores": "REAL",
        "silicon_cores": "REAL",
        "vcf_licensable_cores": "REAL",
        "vcf_license_core_basis": "TEXT",
        "logical_threads": "REAL",
        "cpu_vendor": "TEXT",
        "cpu_model": "TEXT",
        "memory_gib": "REAL",
        "raw_storage_tb": "REAL",
        "raw_storage_tb_osa": "REAL",
        "raw_storage_tb_esa": "REAL",
        "vsan_architecture": "TEXT",
        "storage_only": "INTEGER",
        "availability": "TEXT",
        "catalog_reviewed": "TEXT",
    },
    "sizing_summary": {
        "target": "TEXT",
        "policy": "TEXT",
        "policy_name": "TEXT",
        "topology": "TEXT",
        "status": "TEXT",
        "source_cluster_count": "INTEGER",
        "target_cluster_count": "INTEGER",
        "workload_vms": "INTEGER",
        "compute_vms": "INTEGER",
        "powered_on_workload_vms": "INTEGER",
        "powered_off_or_other_workload_vms": "INTEGER",
        "templates": "INTEGER",
        "configured_vcpus": "REAL",
        "configured_memory_gib": "REAL",
        "provisioned_storage_tib": "REAL",
        "in_use_storage_tib": "REAL",
        "storage_required_tib": "REAL",
        "total_hosts": "INTEGER",
        "vcf_licensable_cores": "REAL",
    },
    "sizing_cluster": {
        "target": "TEXT",
        "policy": "TEXT",
        "topology": "TEXT",
        "status": "TEXT",
        "target_cluster": "TEXT",
        "role": "TEXT",
        "source_clusters": "TEXT",
        "node_type": "TEXT",
        "selection_strategy": "TEXT",
        "vm_count": "INTEGER",
        "vcpus": "REAL",
        "memory_gib": "REAL",
        "provisioned_storage_tib": "REAL",
        "in_use_storage_tib": "REAL",
        "storage_required_tib": "REAL",
        "provider_minimum_hosts": "INTEGER",
        "normal_cpu_floor": "INTEGER",
        "normal_memory_floor": "INTEGER",
        "failure_cpu_floor": "INTEGER",
        "failure_memory_floor": "INTEGER",
        "total_hosts": "INTEGER",
        "binding_constraints": "TEXT",
        "normal_cpu_utilization_percent": "REAL",
        "normal_memory_utilization_percent": "REAL",
        "post_failure_cpu_utilization_percent": "REAL",
        "post_failure_memory_utilization_percent": "REAL",
        "one_host_loss_validated": "INTEGER",
        "vcf_licensable_cores": "REAL",
    },
}

DEFAULT_SELECT = {
    "vm": ("vm", "power_state", "cluster", "host", "guest_os_family", "cpus", "memory_mib"),
    "host": (
        "host",
        "cluster",
        "vendor",
        "model",
        "cpu_vendor",
        "cpu_model",
        "ht_available",
        "ht_active",
        "cores",
        "memory_mib",
        "esxi_version",
    ),
    "cluster": ("cluster", "host_count", "powered_on_vms", "cpu_ratio", "memory_ratio"),
    "datastore": ("datastore", "type", "cluster", "capacity_mib", "provisioned_mib", "free_mib", "free_percent"),
    "license": ("name", "cost_unit", "total", "used", "expiration_date", "vi_sdk_server"),
    "vcf_license": (
        "host",
        "cluster",
        "cpu_sockets",
        "cores_per_cpu",
        "physical_cores",
        "vcf_licensable_cores",
        "core_minimum_adjustment",
    ),
    "vcf_license_summary": (
        "host_count",
        "physical_cores",
        "vcf_licensable_cores",
        "core_calculation_complete",
        "vsan_entitlement_tib",
        "vsan_capacity_tib",
        "vsan_capacity_evidence",
        "vsan_capacity_is_raw",
        "vsan_addon_required_tib",
        "vsan_entitlement_surplus_tib",
    ),
    "disk": ("vm", "power_state", "cluster", "disk", "capacity_mib", "disk_mode", "raw", "sharing_mode"),
    "network": ("vm", "power_state", "network", "switch", "connected", "cluster"),
    "snapshot": ("vm", "cluster", "created_at", "size_mib", "power_state"),
    "dvport": ("port_group", "switch", "vlan", "binding_type"),
    "cdrom": ("vm", "power_state", "cluster", "connected", "starts_connected", "device_type"),
    "usb": ("vm", "power_state", "cluster", "connected", "device_type"),
    "migration_method": ("target", "vm", "cluster", "vcenter", "method", "status", "reason_ids"),
    "migration_finding": (
        "target",
        "finding_id",
        "category",
        "vm",
        "cluster",
        "vcenter",
        "status",
        "methods",
        "summary",
    ),
    "target_node": (
        "target",
        "node_type",
        "shape_series",
        "physical_cores",
        "configured_physical_cores",
        "silicon_cores",
        "vcf_licensable_cores",
        "vcf_license_core_basis",
        "logical_threads",
        "cpu_vendor",
        "cpu_model",
        "memory_gib",
        "raw_storage_tb",
        "raw_storage_tb_osa",
        "raw_storage_tb_esa",
        "storage_only",
        "availability",
    ),
    "sizing_summary": (
        "target",
        "policy",
        "policy_name",
        "topology",
        "status",
        "source_cluster_count",
        "target_cluster_count",
        "compute_vms",
        "configured_vcpus",
        "configured_memory_gib",
        "storage_required_tib",
        "total_hosts",
        "vcf_licensable_cores",
    ),
    "sizing_cluster": (
        "target",
        "policy",
        "topology",
        "target_cluster",
        "role",
        "node_type",
        "vcpus",
        "memory_gib",
        "provider_minimum_hosts",
        "failure_cpu_floor",
        "failure_memory_floor",
        "total_hosts",
        "one_host_loss_validated",
        "vcf_licensable_cores",
    ),
}

TEMPLATE_ENTITIES = {"vm", "disk", "network", "snapshot", "cdrom", "usb"}
WORKLOAD_ONLY_ENTITIES = {"migration_method", "migration_finding"}
FILTER_OPERATORS = {"eq", "ne", "contains", "startswith", "endswith", "gt", "gte", "lt", "lte", "in"}
AGGREGATES = {"sum", "avg", "min", "max", "count_distinct"}


def _sha256(path):
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _text(row, *headers):
    return str(RVTOOLS._value(row, *headers, default="")).strip()


def _number(row, *headers):
    return RVTOOLS._whole_or_float(RVTOOLS._value(row, *headers))


def _optional_number(row, *headers):
    value = RVTOOLS._value(row, *headers)
    return None if value in (None, "") else RVTOOLS._whole_or_float(value)


def _bool(row, *headers):
    return int(RVTOOLS._truthy(RVTOOLS._value(row, *headers)))


def _optional_bool(row, *headers):
    value = RVTOOLS._value(row, *headers)
    status = RVTOOLS._optional_truthy(value)
    return None if status is None else int(status)


def _vcf_core_values(row):
    requirement = host_core_requirement(row, value_getter=RVTOOLS._value)
    return (
        requirement["cpu_sockets"],
        requirement["cores_per_cpu"],
        requirement["physical_silicon_cores"],
        requirement["vcf_licensable_cores"],
        requirement["minimum_core_adjustment"],
    )


def _tib_from_mib(value):
    return round(float(value) / (1024 * 1024), 4)


def _capacity_balance(entitlement_tib, capacity_tib):
    if entitlement_tib is None or capacity_tib is None:
        return None, None
    return (
        round(max(capacity_tib - entitlement_tib, 0), 4),
        round(max(entitlement_tib - capacity_tib, 0), 4),
    )


def _create_table(connection, entity):
    columns = ", ".join(f'"{name}" {kind}' for name, kind in ENTITY_SCHEMAS[entity].items())
    connection.execute(f'CREATE TABLE "{entity}" ({columns})')


def _insert_rows(connection, entity, values):
    columns = tuple(ENTITY_SCHEMAS[entity])
    placeholders = ", ".join("?" for _ in columns)
    column_sql = ", ".join(f'"{column}"' for column in columns)
    connection.executemany(
        f'INSERT INTO "{entity}" ({column_sql}) VALUES ({placeholders})',
        ([row.get(column) for column in columns] for row in values),
    )


def _build_index(connection, workbook_path, digest):
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if "vInfo" not in workbook.sheetnames:
            raise ValueError("not an RVTools workbook: required vInfo sheet is missing")
        sheetnames = set(workbook.sheetnames)
        rows = {sheet: RVTOOLS._read_rows(workbook, sheet) for sheet in INDEXED_SHEETS}
    finally:
        workbook.close()

    connection.execute("CREATE TABLE metadata (key TEXT PRIMARY KEY, value TEXT NOT NULL)")
    connection.executemany(
        "INSERT INTO metadata (key, value) VALUES (?, ?)",
        (
            ("schema_version", INDEX_SCHEMA_VERSION),
            ("source_sha256", digest),
            ("source_file", workbook_path.name),
        ),
    )
    connection.execute(
        "CREATE TABLE query_warning (code TEXT, scope TEXT, message TEXT)"
    )
    for entity in ENTITY_SCHEMAS:
        _create_table(connection, entity)

    tool_status = {
        _text(row, "VM"): _text(row, "Tools")
        for row in rows["vTools"]
        if _text(row, "VM")
    }
    prepared_sizing_input = RVTOOLS.prepare_sizing_input(rows)
    normalized_inventory = prepared_sizing_input["inventory"]
    vm_values = []
    for row, normalized in zip(rows["vInfo"], normalized_inventory):
        guest_os, guest_os_source = RVTOOLS._guest_os(row)
        vm_name = _text(row, "VM") or "(unnamed VM)"
        vm_values.append(
            {
                "vm": vm_name,
                "power_state": _text(row, "Powerstate"),
                "template": _bool(row, "Template"),
                "cluster": normalized["cluster"],
                "host": _text(row, "Host"),
                "guest_os": guest_os,
                "guest_os_family": RVTOOLS._guest_os_family(guest_os),
                "guest_os_source": guest_os_source,
                "cpus": _number(row, "CPUs"),
                "memory_mib": _number(row, "Memory", "Size MiB"),
                "provisioned_mib": _number(row, "Provisioned MiB"),
                "in_use_mib": _number(row, "In Use MiB"),
                "hardware_version": _text(row, "HW version"),
                "tools_status": tool_status.get(vm_name, ""),
                "consolidation_needed": _bool(row, "Consolidation Needed"),
            }
        )
    _insert_rows(connection, "vm", vm_values)
    vm_context = {row["vm"]: row for row in vm_values}

    _insert_rows(
        connection,
        "host",
        (
            {
                "host": _text(row, "Host") or "(unnamed host)",
                "cluster": _text(row, "Cluster") or "(unassigned)",
                "vendor": _text(row, "Vendor"),
                "model": _text(row, "Model"),
                "cpu_vendor": RVTOOLS._cpu_vendor(_text(row, "CPU Model")),
                "cpu_model": _text(row, "CPU Model"),
                "ht_available": _optional_bool(row, "HT Available"),
                "ht_active": _optional_bool(row, "HT Active"),
                "cpu_sockets": _number(row, "# CPU"),
                "cores_per_cpu": _number(row, "Cores per CPU"),
                "esxi_version": _text(row, "ESX Version"),
                "cores": _number(row, "# Cores", "Cores"),
                "cpu_speed_mhz": _number(row, "Speed"),
                "memory_mib": _number(row, "# Memory", "Memory"),
                "cpu_usage_percent": _number(row, "CPU usage %"),
                "memory_usage_percent": _number(row, "Memory usage %"),
                "bios_vendor": _text(row, "BIOS Vendor"),
                "bios_version": _text(row, "BIOS Version"),
                "bios_date": _text(row, "BIOS Date"),
            }
            for row in rows["vHost"]
        ),
    )

    _insert_rows(
        connection,
        "license",
        (
            {
                "name": _text(row, "Name"),
                "cost_unit": _text(row, "Cost Unit"),
                "total": _optional_number(row, "Total"),
                "used": _optional_number(row, "Used"),
                "expiration_date": RVTOOLS._serializable(
                    RVTOOLS._value(row, "Expiration Date")
                ),
                "vi_sdk_server": _text(row, "VI SDK Server"),
            }
            for row in rows["vLicense"]
        ),
    )

    vcf_host_values = []
    for row in rows["vHost"]:
        sockets, cores_per_cpu, physical_cores, licensable_cores, adjustment = (
            _vcf_core_values(row)
        )
        vcf_host_values.append(
            {
                "host": _text(row, "Host") or "(unnamed host)",
                "cluster": _text(row, "Cluster") or "(unassigned)",
                "cpu_sockets": sockets,
                "cores_per_cpu": cores_per_cpu,
                "physical_cores": physical_cores,
                "vcf_licensable_cores": licensable_cores,
                "core_minimum_adjustment": adjustment,
            }
        )
    _insert_rows(connection, "vcf_license", vcf_host_values)

    core_calculation_complete = bool(vcf_host_values) and all(
        row["vcf_licensable_cores"] is not None for row in vcf_host_values
    )
    total_physical_cores = float(
        sum(row["physical_cores"] for row in vcf_host_values)
    )
    total_licensable_cores = (
        float(sum(row["vcf_licensable_cores"] for row in vcf_host_values))
        if core_calculation_complete
        else None
    )
    vsan_datastores = [
        row
        for row in rows["vDatastore"]
        if "vsan" in _text(row, "Type").casefold()
        or "vsan:" in _text(row, "URL").casefold()
    ]
    vsan_capacity_mib = sum(_number(row, "Capacity MiB") for row in vsan_datastores)
    vsan_capacity_tib = _tib_from_mib(vsan_capacity_mib) if vsan_capacity_mib > 0 else None
    vsan_evidence = (
        "rvtools_vsan_datastore_capacity_proxy"
        if vsan_capacity_tib is not None
        else "unavailable"
    )
    vsan_addon_tib, vsan_surplus_tib = _capacity_balance(
        total_licensable_cores, vsan_capacity_tib
    )
    _insert_rows(
        connection,
        "vcf_license_summary",
        [
            {
                "host_count": len(vcf_host_values),
                "physical_cores": total_physical_cores,
                "vcf_licensable_cores": total_licensable_cores,
                "core_calculation_complete": int(core_calculation_complete),
                "vsan_entitlement_tib": total_licensable_cores,
                "vsan_capacity_tib": vsan_capacity_tib,
                "vsan_capacity_evidence": vsan_evidence,
                "vsan_capacity_is_raw": 0,
                "vsan_addon_required_tib": vsan_addon_tib,
                "vsan_entitlement_surplus_tib": vsan_surplus_tib,
            }
        ],
    )

    licensing_warnings = []
    if "vLicense" not in sheetnames:
        licensing_warnings.append(
            (
                "missing_vlicense_sheet",
                "license",
                "vLicense is absent; current VMware licence assignments and consumption are unavailable.",
            )
        )
    licensing_warnings.extend(
        (
            "missing_vcf_cpu_topology",
            row["host"],
            "VCF core capacity is unavailable because the host CPU socket/core topology is incomplete.",
        )
        for row in vcf_host_values
        if row["vcf_licensable_cores"] is None
    )
    if vsan_capacity_tib is None:
        licensing_warnings.append(
            (
                "vsan_raw_capacity_unavailable",
                "vcf_license_summary",
                "Raw vSAN capacity is not available from this RVTools export; provide verified raw TiB to calculate add-on or surplus capacity.",
            )
        )
    else:
        licensing_warnings.append(
            (
                "vsan_capacity_proxy",
                "vcf_license_summary",
                "RVTools vSAN datastore capacity is a planning proxy, not confirmed raw physical capacity; provide verified raw TiB for a licensing conclusion.",
            )
        )
    connection.executemany(
        "INSERT INTO query_warning (code, scope, message) VALUES (?, ?, ?)",
        licensing_warnings,
    )

    overcommit = RVTOOLS._overcommit_summary(rows)
    _insert_rows(connection, "cluster", overcommit["clusters"])
    connection.executemany(
        "INSERT INTO query_warning (code, scope, message) VALUES (?, ?, ?)",
        (
            (
                warning["code"],
                str(warning.get("cluster") or warning.get("sheet") or "overcommit"),
                warning["message"],
            )
            for warning in overcommit["warnings"]
        ),
    )

    _insert_rows(
        connection,
        "datastore",
        (
            {
                "datastore": _text(row, "Name") or "(unnamed datastore)",
                "type": _text(row, "Type"),
                "cluster": _text(row, "Cluster name", "Cluster") or "(unassigned)",
                "capacity_mib": _number(row, "Capacity MiB"),
                "provisioned_mib": _number(row, "Provisioned MiB"),
                "in_use_mib": _number(row, "In Use MiB"),
                "free_mib": _number(row, "Free MiB"),
                "free_percent": _number(row, "Free %"),
                "accessible": _bool(row, "Accessible"),
            }
            for row in rows["vDatastore"]
        ),
    )
    _insert_rows(
        connection,
        "disk",
        (
            {
                "vm": _text(row, "VM"),
                "template": _bool(row, "Template"),
                "power_state": vm_context.get(_text(row, "VM"), {}).get("power_state", ""),
                "cluster": vm_context.get(_text(row, "VM"), {}).get("cluster", ""),
                "host": vm_context.get(_text(row, "VM"), {}).get("host", ""),
                "disk": _text(row, "Disk"),
                "capacity_mib": _number(row, "Capacity MiB"),
                "raw": _bool(row, "Raw"),
                "disk_mode": _text(row, "Disk Mode"),
                "sharing_mode": _text(row, "Sharing mode"),
                "raw_compatibility_mode": _text(row, "Raw Comp. Mode"),
            }
            for row in rows["vDisk"]
        ),
    )
    _insert_rows(
        connection,
        "network",
        (
            {
                "vm": _text(row, "VM"),
                "template": _bool(row, "Template"),
                "power_state": vm_context.get(_text(row, "VM"), {}).get("power_state", ""),
                "network": _text(row, "Network"),
                "switch": _text(row, "Switch"),
                "connected": _bool(row, "Connected"),
                "cluster": _text(row, "Cluster")
                or vm_context.get(_text(row, "VM"), {}).get("cluster", ""),
                "host": _text(row, "Host")
                or vm_context.get(_text(row, "VM"), {}).get("host", ""),
            }
            for row in rows["vNetwork"]
        ),
    )
    _insert_rows(
        connection,
        "snapshot",
        (
            {
                "vm": _text(row, "VM"),
                "template": _bool(row, "Template"),
                "cluster": vm_context.get(_text(row, "VM"), {}).get("cluster", ""),
                "host": vm_context.get(_text(row, "VM"), {}).get("host", ""),
                "created_at": RVTOOLS._serializable(RVTOOLS._value(row, "Date / time")),
                "size_mib": _number(row, "Size MiB (total)", "Size MiB (vmsn)"),
                "power_state": _text(row, "Powerstate")
                or vm_context.get(_text(row, "VM"), {}).get("power_state", ""),
            }
            for row in rows["vSnapshot"]
        ),
    )
    _insert_rows(
        connection,
        "dvport",
        (
            {
                "port_group": _text(row, "Port"),
                "switch": _text(row, "Switch"),
                "vlan": _text(row, "VLAN"),
                "allow_promiscuous": _bool(row, "Allow Promiscuous"),
                "mac_changes": _bool(row, "Mac Changes"),
                "forged_transmits": _bool(row, "Forged Transmits"),
                "binding_type": _text(row, "Type"),
            }
            for row in rows["dvPort"]
        ),
    )
    _insert_rows(
        connection,
        "cdrom",
        (
            {
                "vm": _text(row, "VM"),
                "template": _bool(row, "Template"),
                "power_state": _text(row, "Powerstate")
                or vm_context.get(_text(row, "VM"), {}).get("power_state", ""),
                "cluster": vm_context.get(_text(row, "VM"), {}).get("cluster", ""),
                "host": vm_context.get(_text(row, "VM"), {}).get("host", ""),
                "connected": _bool(row, "Connected"),
                "starts_connected": _bool(row, "Starts Connected"),
                "device_type": _text(row, "Device Type"),
            }
            for row in rows["vCD"]
        ),
    )
    _insert_rows(
        connection,
        "usb",
        (
            {
                "vm": _text(row, "VM"),
                "template": _bool(row, "Template"),
                "power_state": _text(row, "Powerstate")
                or vm_context.get(_text(row, "VM"), {}).get("power_state", ""),
                "cluster": vm_context.get(_text(row, "VM"), {}).get("cluster", ""),
                "host": vm_context.get(_text(row, "VM"), {}).get("host", ""),
                "connected": _bool(row, "Connected"),
                "device_type": _text(row, "Device Type"),
            }
            for row in rows["vUSB"]
        ),
    )
    migration_methods = []
    migration_findings = []
    target_nodes = []
    for target_id in RVTOOLS.TARGET_IDS:
        assessment = RVTOOLS.assess_migration(rows, target_id)
        for item in assessment["vm_methods"]:
            for method, outcome in item["methods"].items():
                migration_methods.append(
                    {
                        "target": target_id,
                        "vm": item["vm"],
                        "cluster": item["cluster"],
                        "host": item["host"],
                        "vcenter": item["vcenter"],
                        "method": method,
                        "status": outcome["status"],
                        "reason_ids": ",".join(outcome["reason_ids"]),
                    }
                )
        for finding in assessment["findings"]:
            migration_findings.append(
                {
                    **finding,
                    "methods": ",".join(finding["methods"]),
                }
            )
        for node in assessment["target"]["nodes"]:
            target_nodes.append(
                {
                    "target": target_id,
                    "node_type": node["id"],
                    "cpu_vendor": assessment["target"]["cpu_vendor"],
                    **{key: value for key, value in node.items() if key != "id"},
                }
            )
    _insert_rows(connection, "migration_method", migration_methods)
    _insert_rows(connection, "migration_finding", migration_findings)
    _insert_rows(connection, "target_node", target_nodes)

    sizing_summaries = []
    sizing_clusters = []
    sizing_errors = set()
    sizing_runtime_warnings = set()
    for target_id in RVTOOLS.TARGET_IDS:
        for policy_id in RVTOOLS.POLICIES:
            for topology in ("consolidated", "source_aligned"):
                try:
                    sizing = RVTOOLS.size_environment(
                        rows,
                        target_id,
                        policy_id=policy_id,
                        topology=topology,
                        prepared=prepared_sizing_input,
                    )
                except ValueError as exc:
                    sizing_errors.add(str(exc))
                    continue
                for warning in sizing["warnings"]:
                    message = warning["message"]
                    if warning.get("count") is not None:
                        message = f"{message} Affected inventory objects: {warning['count']}."
                    sizing_runtime_warnings.add((warning["code"], message))
                totals = sizing["totals"]
                sizing_summaries.append(
                    {
                        "target": target_id,
                        "policy": policy_id,
                        "policy_name": sizing["policy"]["display_name"],
                        "topology": topology,
                        "status": sizing["status"],
                        "source_cluster_count": totals["source_clusters"],
                        "target_cluster_count": totals["target_clusters"],
                        "workload_vms": totals["workload_vms"],
                        "compute_vms": totals["compute_vms"],
                        "powered_on_workload_vms": totals["powered_on_workload_vms"],
                        "powered_off_or_other_workload_vms": totals[
                            "powered_off_or_other_workload_vms"
                        ],
                        "templates": totals["templates"],
                        "configured_vcpus": totals["configured_vcpus"],
                        "configured_memory_gib": totals["configured_memory_gib"],
                        "provisioned_storage_tib": totals["provisioned_storage_tib"],
                        "in_use_storage_tib": totals["in_use_storage_tib"],
                        "storage_required_tib": totals["storage_required_tib"],
                        "total_hosts": totals["total_hosts"],
                        "vcf_licensable_cores": totals["vcf_licensable_cores"],
                    }
                )
                for cluster in sizing["clusters"]:
                    demand = cluster["demand"]
                    recommendation = cluster["recommendation"] or {}
                    sizing_clusters.append(
                        {
                            "target": target_id,
                            "policy": policy_id,
                            "topology": topology,
                            "status": cluster["status"],
                            "target_cluster": cluster["target_cluster"],
                            "role": cluster["role"],
                            "source_clusters": ",".join(cluster["source_clusters"]),
                            "node_type": recommendation.get("node_type"),
                            "selection_strategy": recommendation.get(
                                "selection_strategy"
                            ),
                            "vm_count": demand["vm_count"],
                            "vcpus": demand["vcpus"],
                            "memory_gib": demand["memory_gib"],
                            "provisioned_storage_tib": demand[
                                "provisioned_storage_tib"
                            ],
                            "in_use_storage_tib": demand["in_use_storage_tib"],
                            "storage_required_tib": demand["storage_required_tib"],
                            "provider_minimum_hosts": recommendation.get(
                                "provider_minimum_hosts"
                            ),
                            "normal_cpu_floor": recommendation.get("normal_cpu_floor"),
                            "normal_memory_floor": recommendation.get(
                                "normal_memory_floor"
                            ),
                            "failure_cpu_floor": recommendation.get(
                                "failure_cpu_floor"
                            ),
                            "failure_memory_floor": recommendation.get(
                                "failure_memory_floor"
                            ),
                            "total_hosts": recommendation.get("total_hosts"),
                            "binding_constraints": ",".join(
                                recommendation.get("binding_constraints", [])
                            ),
                            "normal_cpu_utilization_percent": recommendation.get(
                                "normal_cpu_utilization_percent"
                            ),
                            "normal_memory_utilization_percent": recommendation.get(
                                "normal_memory_utilization_percent"
                            ),
                            "post_failure_cpu_utilization_percent": recommendation.get(
                                "post_failure_cpu_utilization_percent"
                            ),
                            "post_failure_memory_utilization_percent": recommendation.get(
                                "post_failure_memory_utilization_percent"
                            ),
                            "one_host_loss_validated": (
                                int(recommendation["one_host_loss_validated"])
                                if "one_host_loss_validated" in recommendation
                                else None
                            ),
                            "vcf_licensable_cores": recommendation.get(
                                "vcf_licensable_cores"
                            ),
                        }
                    )
    _insert_rows(connection, "sizing_summary", sizing_summaries)
    _insert_rows(connection, "sizing_cluster", sizing_clusters)
    connection.executemany(
        "INSERT INTO query_warning (code, scope, message) VALUES (?, ?, ?)",
        (
            (
                "migration_screening_not_validation",
                entity,
                "RVTools provides a planning screen, not a successful HCX Validate result; complete the listed manual gates before execution.",
            )
            for entity in ("migration_method", "migration_finding", "target_node")
        ),
    )
    connection.executemany(
        "INSERT INTO query_warning (code, scope, message) VALUES (?, ?, ?)",
        (
            (code, entity, message)
            for entity in ("sizing_summary", "sizing_cluster")
            for code, message in sorted(sizing_runtime_warnings)
            if code != "configuration_only_sizing"
        ),
    )
    connection.executemany(
        "INSERT INTO query_warning (code, scope, message) VALUES (?, ?, ?)",
        (
            ("sizing_unavailable", entity, message)
            for entity in ("sizing_summary", "sizing_cluster")
            for message in sorted(sizing_errors)
        ),
    )
    connection.executemany(
        "INSERT INTO query_warning (code, scope, message) VALUES (?, ?, ?)",
        (
            (
                "configuration_only_sizing",
                entity,
                "RVTools configuration is a planning baseline; validate sustained CPU, memory, storage, and network demand before purchase.",
            )
            for entity in ("sizing_summary", "sizing_cluster")
        ),
    )
    connection.commit()


def _cached_connection(workbook_path, index_path):
    digest = _sha256(workbook_path)
    if index_path is None:
        connection = sqlite3.connect(":memory:")
        _build_index(connection, workbook_path, digest)
        return connection, False, digest

    index_path = Path(index_path)
    if index_path.resolve() == workbook_path.resolve():
        raise ValueError("index path must not overwrite the source workbook")
    if index_path.exists():
        try:
            connection = sqlite3.connect(index_path)
            metadata = dict(connection.execute("SELECT key, value FROM metadata"))
            if (
                metadata.get("schema_version") == INDEX_SCHEMA_VERSION
                and metadata.get("source_sha256") == digest
            ):
                return connection, True, digest
            connection.close()
        except sqlite3.Error:
            try:
                connection.close()
            except UnboundLocalError:
                pass

    index_path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{index_path.name}.", suffix=".tmp", dir=index_path.parent
    )
    os.close(descriptor)
    temporary_path = Path(temporary_name)
    try:
        connection = sqlite3.connect(temporary_path)
        _build_index(connection, workbook_path, digest)
        connection.close()
        os.chmod(temporary_path, 0o600)
        os.replace(temporary_path, index_path)
        os.chmod(index_path, 0o600)
    finally:
        if temporary_path.exists():
            temporary_path.unlink()
    return sqlite3.connect(index_path), False, digest


def _validate_field(entity, field):
    if field not in ENTITY_SCHEMAS[entity]:
        raise ValueError(f"field '{field}' is not allowed for entity '{entity}'")
    return field


def _coerce_value(entity, field, value):
    kind = ENTITY_SCHEMAS[entity][field]
    if kind == "INTEGER":
        lowered = str(value).strip().casefold()
        if lowered in {"true", "yes", "on"}:
            return 1
        if lowered in {"false", "no", "off"}:
            return 0
        return int(value)
    if kind == "REAL":
        return float(value)
    return str(value)


def _filter_sql(entity, expression):
    if "=" not in expression:
        raise ValueError(f"filter '{expression}' must use field=value syntax")
    left, value = expression.split("=", 1)
    if "__" in left:
        field, operator = left.rsplit("__", 1)
    else:
        field, operator = left, "eq"
    _validate_field(entity, field)
    if operator not in FILTER_OPERATORS:
        raise ValueError(f"filter operator '{operator}' is not allowed")
    column = f'"{field}"'
    kind = ENTITY_SCHEMAS[entity][field]
    if operator == "in":
        values = [_coerce_value(entity, field, item.strip()) for item in value.split(",")]
        if not values:
            raise ValueError("in filter requires at least one value")
        placeholders = ", ".join("?" for _ in values)
        return f"{column} IN ({placeholders})", values
    if operator in {"contains", "startswith", "endswith"}:
        if kind != "TEXT":
            raise ValueError(f"operator '{operator}' requires a text field")
        pattern = value
        if operator in {"contains", "endswith"}:
            pattern = "%" + pattern
        if operator in {"contains", "startswith"}:
            pattern = pattern + "%"
        return f"LOWER({column}) LIKE LOWER(?)", [pattern]
    sql_operator = {
        "eq": "=",
        "ne": "!=",
        "gt": ">",
        "gte": ">=",
        "lt": "<",
        "lte": "<=",
    }[operator]
    parameter = _coerce_value(entity, field, value)
    if kind == "TEXT" and operator in {"eq", "ne"}:
        return f"{column} {sql_operator} ? COLLATE NOCASE", [parameter]
    return f"{column} {sql_operator} ?", [parameter]


def _metric_sql(entity, metric):
    if metric == "count":
        return "COUNT(*)", "count"
    if ":" not in metric:
        raise ValueError(f"metric '{metric}' is not allowed")
    aggregate, field = metric.split(":", 1)
    if aggregate not in AGGREGATES:
        raise ValueError(f"aggregate '{aggregate}' is not allowed")
    _validate_field(entity, field)
    if aggregate in {"sum", "avg"} and ENTITY_SCHEMAS[entity][field] not in {"INTEGER", "REAL"}:
        raise ValueError(f"aggregate '{aggregate}' requires a numeric field")
    function = "COUNT" if aggregate == "count_distinct" else aggregate.upper()
    distinct = "DISTINCT " if aggregate == "count_distinct" else ""
    alias = f"{aggregate}_{field}"
    return f'{function}({distinct}"{field}")', alias


def _compile_query(plan):
    entity = str(plan.get("entity", "")).strip().casefold()
    if entity not in ENTITY_SCHEMAS:
        raise ValueError(f"entity '{entity}' is not allowed")
    metrics = list(plan.get("metrics") or [])
    group_by = list(plan.get("group_by") or [])
    selected = list(plan.get("select") or [])
    for field in group_by + selected:
        _validate_field(entity, field)
    if metrics and selected and set(selected) - set(group_by):
        raise ValueError("selected fields in an aggregate query must also be group_by fields")
    if metrics:
        output_fields = group_by[:]
        select_parts = [f'"{field}"' for field in group_by]
        for metric in metrics:
            expression, alias = _metric_sql(entity, metric)
            select_parts.append(f'{expression} AS "{alias}"')
            output_fields.append(alias)
    else:
        output_fields = selected or list(DEFAULT_SELECT[entity])
        select_parts = [f'"{field}"' for field in output_fields]

    where_parts = []
    parameters = []
    template_filter_applied = entity in TEMPLATE_ENTITIES and not bool(plan.get("include_templates"))
    templates_excluded = template_filter_applied or entity in WORKLOAD_ONLY_ENTITIES
    if template_filter_applied:
        where_parts.append('"template" = 0')
    for expression in plan.get("filters") or []:
        sql, values = _filter_sql(entity, str(expression))
        where_parts.append(sql)
        parameters.extend(values)

    sql = f'SELECT {", ".join(select_parts)} FROM "{entity}"'
    if where_parts:
        sql += " WHERE " + " AND ".join(where_parts)
    if group_by:
        sql += " GROUP BY " + ", ".join(f'"{field}"' for field in group_by)

    allowed_order_fields = set(output_fields)
    order_parts = []
    for item in plan.get("order_by") or []:
        value = str(item)
        field, _, direction = value.partition(":")
        direction = (direction or "asc").casefold()
        if field not in allowed_order_fields or direction not in {"asc", "desc"}:
            raise ValueError(f"order_by '{item}' is not allowed")
        order_parts.append(f'"{field}" {direction.upper()}')
    if not order_parts and group_by:
        order_parts = [f'"{field}" ASC' for field in group_by]
    if order_parts:
        sql += " ORDER BY " + ", ".join(order_parts)

    limit = int(plan.get("limit", 25))
    if limit < 1 or limit > 100:
        raise ValueError("limit must be between 1 and 100")
    sql += " LIMIT ?"
    parameters.append(limit + 1)
    return entity, sql, parameters, output_fields, templates_excluded, limit


def run_query(workbook, plan, *, index_path=None):
    workbook_path = Path(workbook)
    if not workbook_path.exists():
        raise FileNotFoundError(workbook_path)
    if workbook_path.suffix.casefold() != ".xlsx":
        raise ValueError("RVTools query input must be an .xlsx workbook")
    vsan_raw_tib = plan.get("vsan_raw_tib")
    if vsan_raw_tib is not None:
        if str(plan.get("entity", "")).strip().casefold() != "vcf_license_summary":
            raise ValueError("vsan_raw_tib is only allowed for entity 'vcf_license_summary'")
        if plan.get("select") or plan.get("metrics") or plan.get("filters") or plan.get("group_by"):
            raise ValueError("vsan_raw_tib requires the unfiltered default VCF licence summary")
        vsan_raw_tib = float(vsan_raw_tib)
        if not math.isfinite(vsan_raw_tib) or vsan_raw_tib < 0:
            raise ValueError("vsan_raw_tib must be a finite non-negative number")
    entity, sql, parameters, fields, templates_excluded, limit = _compile_query(plan)
    connection, index_reused, digest = _cached_connection(workbook_path, index_path)
    connection.row_factory = sqlite3.Row
    try:
        fetched = connection.execute(sql, parameters).fetchall()
        cluster_scoped = entity == "cluster" or "cluster" in (plan.get("group_by") or [])
        exact_cluster = None
        for query_filter in plan.get("filters") or []:
            filter_text = str(query_filter)
            if filter_text.casefold().startswith("cluster="):
                exact_cluster = filter_text.split("=", 1)[1]
            elif filter_text.casefold().startswith("cluster__eq="):
                exact_cluster = filter_text.split("=", 1)[1]
        if entity == "cluster":
            if exact_cluster is not None:
                warning_query = (
                    "SELECT code, scope, message FROM query_warning "
                    "WHERE scope = ? COLLATE NOCASE OR code = 'multi_vcenter_sources'"
                )
                warning_parameters = (exact_cluster,)
            else:
                warning_query = (
                    "SELECT code, scope, message FROM query_warning "
                    "WHERE code IN ('missing_physical_cpu_capacity', "
                    "'missing_physical_memory_capacity', 'multi_vcenter_sources')"
                )
                warning_parameters = ()
        elif entity == "vcf_license":
            warning_query = (
                "SELECT code, scope, message FROM query_warning "
                "WHERE code IN ('missing_vcf_cpu_topology', 'multi_vcenter_sources')"
            )
            warning_parameters = ()
        elif entity == "vcf_license_summary":
            warning_query = (
                "SELECT code, scope, message FROM query_warning "
                "WHERE code IN ('missing_vcf_cpu_topology', 'vsan_capacity_proxy', "
                "'vsan_raw_capacity_unavailable', 'multi_vcenter_sources')"
            )
            warning_parameters = ()
        elif entity == "license":
            warning_query = (
                "SELECT code, scope, message FROM query_warning "
                "WHERE code = 'missing_vlicense_sheet'"
            )
            warning_parameters = ()
        elif entity in {"migration_method", "migration_finding", "target_node"}:
            warning_query = (
                "SELECT code, scope, message FROM query_warning "
                "WHERE code = 'migration_screening_not_validation' AND scope = ?"
            )
            warning_parameters = (entity,)
        elif entity in {"sizing_summary", "sizing_cluster"}:
            warning_query = (
                "SELECT code, scope, message FROM query_warning "
                "WHERE scope = ?"
            )
            warning_parameters = (entity,)
        elif cluster_scoped:
            warning_query = (
                "SELECT code, scope, message FROM query_warning "
                "WHERE code = 'multi_vcenter_sources'"
            )
            warning_parameters = ()
        else:
            warning_query = "SELECT code, scope, message FROM query_warning WHERE 0"
            warning_parameters = ()
        warnings = [
            dict(row)
            for row in connection.execute(warning_query, warning_parameters)
        ]
    finally:
        connection.close()
    truncated = len(fetched) > limit
    rows = [dict(row) for row in fetched[:limit]]
    if vsan_raw_tib is not None and rows:
        entitlement_tib = rows[0]["vsan_entitlement_tib"]
        addon_tib, surplus_tib = _capacity_balance(entitlement_tib, vsan_raw_tib)
        rows[0].update(
            {
                "vsan_capacity_tib": vsan_raw_tib,
                "vsan_capacity_evidence": "user_supplied_raw_tib",
                "vsan_capacity_is_raw": 1,
                "vsan_addon_required_tib": addon_tib,
                "vsan_entitlement_surplus_tib": surplus_tib,
            }
        )
        warnings = [
            warning
            for warning in warnings
            if warning["code"] not in {"vsan_capacity_proxy", "vsan_raw_capacity_unavailable"}
        ]
    return {
        "source": {"file": workbook_path.name, "sha256": digest},
        "query": {
            "entity": entity,
            "select": fields,
            "filters": list(plan.get("filters") or []),
            "group_by": list(plan.get("group_by") or []),
            "metrics": list(plan.get("metrics") or []),
            "vsan_raw_tib": vsan_raw_tib,
        },
        "scope": {
            "templates_excluded": templates_excluded,
            "row_limit": limit,
        },
        "rows": rows,
        "row_count": len(rows),
        "truncated": truncated,
        "warnings": warnings,
        "index_reused": index_reused,
    }


def main(argv=None):
    argument_parser = argparse.ArgumentParser(
        description="Query an RVTools workbook through a local allowlisted SQLite index."
    )
    argument_parser.add_argument("workbook", type=Path, help="Path to the RVTools .xlsx export")
    argument_parser.add_argument("--entity", required=True, choices=sorted(ENTITY_SCHEMAS))
    argument_parser.add_argument("--select", action="append", default=[], help="Field to return; repeat as needed")
    argument_parser.add_argument("--metric", action="append", default=[], help="count or aggregate:field; repeat as needed")
    argument_parser.add_argument("--filter", action="append", default=[], help="field[__operator]=value; repeat as needed")
    argument_parser.add_argument("--group-by", action="append", default=[], help="Grouping field; repeat as needed")
    argument_parser.add_argument("--order-by", action="append", default=[], help="Output field[:asc|desc]; repeat as needed")
    argument_parser.add_argument("--limit", type=int, default=25, help="Maximum rows returned, from 1 to 100")
    argument_parser.add_argument("--include-templates", action="store_true", help="Include templates for VM-linked entities")
    argument_parser.add_argument("--index", type=Path, help="Optional reusable SQLite index path")
    argument_parser.add_argument(
        "--vsan-raw-tib",
        type=float,
        help="Verified raw vSAN capacity in TiB for the VCF licence summary",
    )
    argument_parser.add_argument("--pretty", action="store_true", help="Indent JSON output")
    args = argument_parser.parse_args(argv)
    plan = {
        "entity": args.entity,
        "select": args.select,
        "metrics": args.metric,
        "filters": args.filter,
        "group_by": args.group_by,
        "order_by": args.order_by,
        "limit": args.limit,
        "include_templates": args.include_templates,
        "vsan_raw_tib": args.vsan_raw_tib,
    }
    try:
        result = run_query(args.workbook, plan, index_path=args.index)
    except (FileNotFoundError, OSError, ValueError, sqlite3.Error) as exc:
        argument_parser.exit(2, f"error: {exc}\n")
    json.dump(result, sys.stdout, indent=2 if args.pretty else None, sort_keys=args.pretty)
    sys.stdout.write("\n")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
