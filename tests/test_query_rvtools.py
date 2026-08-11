import importlib.util
import json
import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

from openpyxl import Workbook, load_workbook


MODULE_PATH = Path(__file__).parents[1] / "scripts" / "query_rvtools.py"
SPEC = importlib.util.spec_from_file_location("query_rvtools", MODULE_PATH)
QUERY = importlib.util.module_from_spec(SPEC)
assert SPEC.loader is not None
SPEC.loader.exec_module(QUERY)
TARGETS = importlib.import_module("rvtools.targets")


def add_sheet(workbook, name, headers, rows):
    worksheet = workbook.create_sheet(name)
    worksheet.append(headers)
    for row in rows:
        worksheet.append(row)


class QueryRVToolsTests(unittest.TestCase):
    def setUp(self):
        self.tempdir = tempfile.TemporaryDirectory()
        self.path = Path(self.tempdir.name) / "query-rvtools.xlsx"
        workbook = Workbook()
        workbook.remove(workbook.active)
        add_sheet(
            workbook,
            "vInfo",
            [
                "VM",
                "Powerstate",
                "Template",
                "CPUs",
                "Memory",
                "Provisioned MiB",
                "In Use MiB",
                "Cluster",
                "Host",
                "HW version",
                "Consolidation Needed",
                "OS according to the configuration file",
                "OS according to the VMware Tools",
            ],
            [
                ["linux-on", "poweredOn", False, 4, 8192, 102400, 51200, "cluster-a", "esx-01", "vmx-19", False, "Other Linux", "Ubuntu Linux 24.04"],
                ["windows-off", "poweredOff", False, 2, 4096, 51200, 20480, "cluster-a", "esx-02", "vmx-15", False, "Microsoft Windows Server 2022", ""],
                ["bsd-on", "poweredOn", False, 2, 2048, 20480, 10240, "cluster-b", "esx-03", "vmx-19", False, "FreeBSD 13", ""],
                ["linux-template", "poweredOff", True, 8, 16384, 204800, 102400, "cluster-a", "esx-01", "vmx-19", False, "Ubuntu Linux", ""],
            ],
        )
        add_sheet(
            workbook,
            "vHost",
            [
                "Host",
                "Cluster",
                "CPU Model",
                "ESX Version",
                "# Cores",
                "# Memory",
                "Vendor",
                "Model",
                "HT Available",
                "HT Active",
                "# CPU",
                "Cores per CPU",
                "Speed",
                "BIOS Vendor",
                "BIOS Version",
                "BIOS Date",
            ],
            [
                ["esx-01", "cluster-a", "Intel(R) Xeon Gold 6248", "8.0.3", 16, 65536, "Cisco Systems Inc", "HXAF240C-M5SX", True, True, 2, 8, 2500, "Cisco Systems, Inc.", "C240M5.4.1.3m", "2022-07-08"],
                ["esx-02", "cluster-a", "AMD EPYC 7543", "8.0.3", 16, 65536, "Dell Inc.", "PowerEdge R7525", True, False, 2, 8, 2800, "Dell Inc.", "2.19.1", "2024-01-15"],
                ["esx-03", "cluster-b", "Intel Xeon", "8.0.3", 8, 32768, "", "", "", "", 1, 8, 2200, "", "", ""],
            ],
        )
        add_sheet(
            workbook,
            "vTools",
            ["VM", "Template", "Tools", "Powerstate"],
            [
                ["linux-on", False, "toolsOk", "poweredOn"],
                ["windows-off", False, "toolsNotRunning", "poweredOff"],
                ["linux-template", True, "toolsNotRunning", "poweredOff"],
            ],
        )
        add_sheet(
            workbook,
            "vDatastore",
            [
                "Name",
                "Capacity MiB",
                "Provisioned MiB",
                "In Use MiB",
                "Free MiB",
                "Free %",
                "Accessible",
                "Type",
                "Cluster name",
                "URL",
            ],
            [
                ["vsan-01", 100 * 1024 * 1024, 700000, 600000, 400000, 40, True, "vsan", "cluster-a", "ds:///vmfs/volumes/vsan:demo"],
                ["san-01", 50 * 1024 * 1024, 700000, 600000, 400000, 40, True, "VMFS", "cluster-b", "ds:///vmfs/volumes/demo"],
            ],
        )
        add_sheet(
            workbook,
            "vLicense",
            ["Name", "Key", "Labels", "Cost Unit", "Total", "Used", "Expiration Date", "Features", "VI SDK Server"],
            [
                ["vSphere 8 Enterprise Plus", "AAAAA-BBBBB-CCCCC-DDDDD-EEEEE", "Production", "cpuPackage", 6, 5, "2027-12-31", "feature-list", "vcenter-a.example"],
                ["vCenter Server 8 Standard", "FFFFF-GGGGG-HHHHH-IIIII-JJJJJ", "Management", "instance", 2, 1, "2027-12-31", "feature-list", "vcenter-a.example"],
            ],
        )
        add_sheet(
            workbook,
            "vDisk",
            ["VM", "Template", "Disk", "Capacity MiB", "Raw", "Disk Mode", "Sharing mode", "Raw Comp. Mode"],
            [["linux-on", False, "Hard disk 1", 102400, False, "persistent", "sharingNone", ""]],
        )
        add_sheet(
            workbook,
            "vNetwork",
            ["VM", "Template", "Network", "Switch", "Connected", "Cluster", "Host"],
            [["linux-on", False, "production", "dvSwitch-01", True, "cluster-a", "esx-01"]],
        )
        add_sheet(
            workbook,
            "vSnapshot",
            ["VM", "Template", "Date / time", "Size MiB (total)", "Powerstate"],
            [["linux-on", False, "2026-01-01T00:00:00", 1024, "poweredOn"]],
        )
        add_sheet(
            workbook,
            "dvPort",
            ["Port", "Switch", "VLAN", "Allow Promiscuous", "Mac Changes", "Forged Transmits", "Type"],
            [["production", "dvSwitch-01", 100, False, False, False, "earlyBinding"]],
        )
        add_sheet(
            workbook,
            "vCD",
            ["VM", "Template", "Powerstate", "Connected", "Starts Connected", "Device Type"],
            [
                ["linux-on", False, "poweredOn", True, True, "ISO image"],
                ["linux-template", True, "poweredOff", True, True, "ISO image"],
            ],
        )
        add_sheet(
            workbook,
            "vUSB",
            ["VM", "Template", "Powerstate", "Connected", "Device Type"],
            [["windows-off", False, "poweredOff", False, "USB device"]],
        )
        workbook.save(self.path)
        workbook.close()

    def tearDown(self):
        self.tempdir.cleanup()

    def test_aggregates_filtered_inventory_and_excludes_templates_by_default(self):
        self.assertTrue(hasattr(QUERY, "run_query"))
        result = QUERY.run_query(
            self.path,
            {
                "entity": "vm",
                "metrics": ["count", "sum:cpus"],
                "filters": ["guest_os_family=linux"],
                "group_by": ["cluster"],
            },
        )

        self.assertEqual(result["rows"], [{"cluster": "cluster-a", "count": 1, "sum_cpus": 4}])
        self.assertTrue(result["scope"]["templates_excluded"])

    def test_query_indexing_closes_source_workbook(self):
        workbook = load_workbook(self.path, read_only=True, data_only=True)
        close_spy = Mock(wraps=workbook.close)
        workbook.close = close_spy

        try:
            with patch.object(QUERY, "load_workbook", return_value=workbook):
                QUERY.run_query(self.path, {"entity": "vm", "metrics": ["count"]})
            close_spy.assert_called_once_with()
        finally:
            if not close_spy.called:
                workbook.close()

    def test_queries_cluster_overcommit_fields(self):
        self.assertTrue(hasattr(QUERY, "run_query"))
        result = QUERY.run_query(
            self.path,
            {
                "entity": "cluster",
                "select": ["cluster", "cpu_ratio", "memory_ratio"],
                "filters": ["cluster=cluster-a"],
            },
        )

        self.assertEqual(
            result["rows"],
            [{"cluster": "cluster-a", "cpu_ratio": 0.12, "memory_ratio": 0.06}],
        )

    def test_queries_host_vendor_model_cpu_topology_hyperthreading_and_bios(self):
        expected_fields = {
            "vendor",
            "model",
            "cpu_vendor",
            "cpu_model",
            "ht_available",
            "ht_active",
            "cpu_sockets",
            "cores_per_cpu",
            "cpu_speed_mhz",
            "bios_vendor",
            "bios_version",
            "bios_date",
        }
        self.assertTrue(expected_fields.issubset(QUERY.ENTITY_SCHEMAS["host"]))

        result = QUERY.run_query(
            self.path,
            {
                "entity": "host",
                "select": ["host", *sorted(expected_fields)],
                "filters": ["host=esx-01"],
            },
        )

        self.assertEqual(
            result["rows"],
            [
                {
                    "host": "esx-01",
                    "bios_date": "2022-07-08",
                    "bios_vendor": "Cisco Systems, Inc.",
                    "bios_version": "C240M5.4.1.3m",
                    "cores_per_cpu": 8.0,
                    "cpu_model": "Intel(R) Xeon Gold 6248",
                    "cpu_sockets": 2.0,
                    "cpu_speed_mhz": 2500.0,
                    "cpu_vendor": "Intel",
                    "ht_active": 1,
                    "ht_available": 1,
                    "model": "HXAF240C-M5SX",
                    "vendor": "Cisco Systems Inc",
                }
            ],
        )

        unknown_ht = QUERY.run_query(
            self.path,
            {
                "entity": "host",
                "select": ["host", "ht_available", "ht_active"],
                "filters": ["host=esx-03"],
            },
        )
        self.assertEqual(
            unknown_ht["rows"],
            [{"host": "esx-03", "ht_available": None, "ht_active": None}],
        )

    def test_filters_vm_linked_entities_by_vm_cluster_context(self):
        self.assertTrue(hasattr(QUERY, "run_query"))
        self.assertIn("cluster", QUERY.ENTITY_SCHEMAS["disk"])
        result = QUERY.run_query(
            self.path,
            {
                "entity": "disk",
                "metrics": ["count", "sum:capacity_mib"],
                "filters": ["cluster=cluster-a", "power_state=poweredOn"],
            },
        )

        self.assertEqual(result["rows"], [{"count": 1, "sum_capacity_mib": 102400}])

    def test_queries_cdrom_and_usb_device_inventory(self):
        self.assertIn("cdrom", QUERY.ENTITY_SCHEMAS)
        self.assertIn("usb", QUERY.ENTITY_SCHEMAS)
        cdrom = QUERY.run_query(
            self.path,
            {"entity": "cdrom", "metrics": ["count"], "filters": ["connected=true"]},
        )
        usb = QUERY.run_query(
            self.path,
            {"entity": "usb", "metrics": ["count"], "filters": ["connected=false"]},
        )

        self.assertEqual(cdrom["rows"], [{"count": 1}])
        self.assertEqual(usb["rows"], [{"count": 1}])

    def test_queries_exact_migration_method_outcomes_by_target(self):
        result = QUERY.run_query(
            self.path,
            {
                "entity": "migration_method",
                "select": ["target", "vm", "method", "status", "reason_ids"],
                "filters": ["target=avs", "vm=linux-on", "method=hcx_vmotion"],
            },
        )

        self.assertEqual(
            result["rows"],
            [
                {
                    "target": "avs",
                    "vm": "linux-on",
                    "method": "hcx_vmotion",
                    "status": "blocked",
                    "reason_ids": "connected_cdrom",
                }
            ],
        )
        self.assertEqual(result["warnings"][0]["code"], "migration_screening_not_validation")
        self.assertTrue(result["scope"]["templates_excluded"])

    def test_queries_provider_node_catalog_and_migration_findings(self):
        nodes = QUERY.run_query(
            self.path,
            {
                "entity": "target_node",
                "select": ["target", "node_type", "physical_cores", "logical_threads"],
                "filters": ["target=gcve", "node_type=ve2-standard-128"],
            },
        )
        findings = QUERY.run_query(
            self.path,
            {
                "entity": "migration_finding",
                "select": ["target", "finding_id", "vm", "status", "methods"],
                "filters": ["target=avs", "finding_id=cross_vendor_cpu"],
            },
        )

        self.assertEqual(
            nodes["rows"],
            [
                {
                    "target": "gcve",
                    "node_type": "ve2-standard-128",
                    "physical_cores": 64.0,
                    "logical_threads": 128.0,
                }
            ],
        )
        self.assertEqual(findings["rows"][0]["vm"], "windows-off")
        self.assertEqual(findings["rows"][0]["status"], "blocked")

    def test_target_catalog_licenses_full_silicon_when_compute_cores_are_reduced(self):
        selected_ocvs = TARGETS.target_profile(
            "ocvs", target_node="BM.Standard.E5.192-96"
        )
        gcve = QUERY.run_query(
            self.path,
            {
                "entity": "target_node",
                "select": [
                    "target",
                    "node_type",
                    "configured_physical_cores",
                    "silicon_cores",
                    "vcf_licensable_cores",
                ],
                "filters": ["target=gcve", "node_type=ve2-standard-64"],
            },
        )
        ocvs = QUERY.run_query(
            self.path,
            {
                "entity": "target_node",
                "select": [
                    "target",
                    "node_type",
                    "shape_series",
                    "cpu_vendor",
                    "configured_physical_cores",
                    "silicon_cores",
                    "vcf_licensable_cores",
                ],
                "filters": ["target=ocvs", "node_type=BM.Standard.E5.192-96"],
            },
        )
        storage_only = QUERY.run_query(
            self.path,
            {
                "entity": "target_node",
                "select": [
                    "node_type",
                    "configured_physical_cores",
                    "silicon_cores",
                    "vcf_licensable_cores",
                    "storage_only",
                ],
                "filters": ["target=gcve", "node_type=ve2-standard-so"],
            },
        )

        self.assertEqual(
            gcve["rows"],
            [
                {
                    "target": "gcve",
                    "node_type": "ve2-standard-64",
                    "configured_physical_cores": 32.0,
                    "silicon_cores": 64.0,
                    "vcf_licensable_cores": 64.0,
                }
            ],
        )
        self.assertEqual(selected_ocvs["cpu_vendor"], "AMD")
        self.assertEqual(selected_ocvs["selected_node"]["vcf_licensable_cores"], 192)
        with self.assertRaisesRegex(ValueError, "uses AMD hosts"):
            TARGETS.target_profile(
                "ocvs",
                target_node="BM.Standard.E5.192-96",
                target_cpu_vendor="Intel",
            )
        self.assertEqual(
            ocvs["rows"],
            [
                {
                    "target": "ocvs",
                    "node_type": "BM.Standard.E5.192-96",
                    "shape_series": "BM.Standard.E5.192",
                    "cpu_vendor": "AMD",
                    "configured_physical_cores": 96.0,
                    "silicon_cores": 192.0,
                    "vcf_licensable_cores": 192.0,
                }
            ],
        )
        self.assertEqual(
            storage_only["rows"],
            [
                {
                    "node_type": "ve2-standard-so",
                    "configured_physical_cores": 0.0,
                    "silicon_cores": 64.0,
                    "vcf_licensable_cores": 64.0,
                    "storage_only": 1,
                }
            ],
        )

    def test_queries_sanitized_current_vmware_license_inventory(self):
        self.assertIn("license", QUERY.ENTITY_SCHEMAS)
        self.assertNotIn("key", QUERY.ENTITY_SCHEMAS["license"])
        self.assertNotIn("features", QUERY.ENTITY_SCHEMAS["license"])
        self.assertNotIn("labels", QUERY.ENTITY_SCHEMAS["license"])

        result = QUERY.run_query(self.path, {"entity": "license", "limit": 10})

        self.assertEqual(
            result["rows"],
            [
                {
                    "name": "vSphere 8 Enterprise Plus",
                    "cost_unit": "cpuPackage",
                    "total": 6.0,
                    "used": 5.0,
                    "expiration_date": "2027-12-31",
                    "vi_sdk_server": "vcenter-a.example",
                },
                {
                    "name": "vCenter Server 8 Standard",
                    "cost_unit": "instance",
                    "total": 2.0,
                    "used": 1.0,
                    "expiration_date": "2027-12-31",
                    "vi_sdk_server": "vcenter-a.example",
                },
            ],
        )
        self.assertNotIn("AAAAA-BBBBB", json.dumps(result))

    def test_calculates_vcf_core_and_vsan_capacity_balance(self):
        hosts = QUERY.run_query(
            self.path,
            {
                "entity": "vcf_license",
                "select": [
                    "host",
                    "cluster",
                    "cpu_sockets",
                    "cores_per_cpu",
                    "physical_cores",
                    "vcf_licensable_cores",
                    "core_minimum_adjustment",
                ],
                "limit": 10,
            },
        )
        self.assertEqual(
            [row["vcf_licensable_cores"] for row in hosts["rows"]],
            [32.0, 32.0, 16.0],
        )
        self.assertEqual(
            [row["core_minimum_adjustment"] for row in hosts["rows"]],
            [16.0, 16.0, 8.0],
        )
        self.assertEqual(hosts["warnings"], [])

        summary = QUERY.run_query(self.path, {"entity": "vcf_license_summary"})
        self.assertEqual(
            summary["rows"],
            [
                {
                    "host_count": 3,
                    "physical_cores": 40.0,
                    "vcf_licensable_cores": 80.0,
                    "core_calculation_complete": 1,
                    "vsan_entitlement_tib": 80.0,
                    "vsan_capacity_tib": 100.0,
                    "vsan_capacity_evidence": "rvtools_vsan_datastore_capacity_proxy",
                    "vsan_capacity_is_raw": 0,
                    "vsan_addon_required_tib": 20.0,
                    "vsan_entitlement_surplus_tib": 0.0,
                }
            ],
        )
        self.assertEqual(
            {warning["code"] for warning in summary["warnings"]},
            {"vsan_capacity_proxy"},
        )

    def test_uses_verified_raw_vsan_capacity_override(self):
        result = QUERY.run_query(
            self.path,
            {"entity": "vcf_license_summary", "vsan_raw_tib": 70.5},
        )

        self.assertEqual(result["rows"][0]["vsan_capacity_tib"], 70.5)
        self.assertEqual(result["rows"][0]["vsan_capacity_evidence"], "user_supplied_raw_tib")
        self.assertEqual(result["rows"][0]["vsan_capacity_is_raw"], 1)
        self.assertEqual(result["rows"][0]["vsan_addon_required_tib"], 0.0)
        self.assertEqual(result["rows"][0]["vsan_entitlement_surplus_tib"], 9.5)
        self.assertEqual(result["warnings"], [])

    def test_requests_raw_vsan_capacity_when_no_vsan_evidence_exists(self):
        workbook = load_workbook(self.path)
        workbook["vDatastore"]["H2"] = "VMFS"
        workbook["vDatastore"]["J2"] = "ds:///vmfs/volumes/demo"
        workbook.save(self.path)
        workbook.close()

        result = QUERY.run_query(self.path, {"entity": "vcf_license_summary"})

        row = result["rows"][0]
        self.assertEqual(row["vsan_capacity_evidence"], "unavailable")
        self.assertIsNone(row["vsan_capacity_tib"])
        self.assertIsNone(row["vsan_addon_required_tib"])
        self.assertIsNone(row["vsan_entitlement_surplus_tib"])
        self.assertIn(
            "vsan_raw_capacity_unavailable",
            {item["code"] for item in result["warnings"]},
        )

    def test_withholds_estate_total_when_host_cpu_topology_is_incomplete(self):
        workbook = load_workbook(self.path)
        workbook["vHost"]["K4"] = None
        workbook["vHost"]["L4"] = None
        workbook.save(self.path)
        workbook.close()

        result = QUERY.run_query(self.path, {"entity": "vcf_license_summary"})

        row = result["rows"][0]
        self.assertEqual(row["core_calculation_complete"], 0)
        self.assertIsNone(row["vcf_licensable_cores"])
        self.assertIsNone(row["vsan_entitlement_tib"])
        self.assertIsNone(row["vsan_addon_required_tib"])
        self.assertIsNone(row["vsan_entitlement_surplus_tib"])
        self.assertIn("missing_vcf_cpu_topology", {item["code"] for item in result["warnings"]})

    def test_rejects_non_allowlisted_fields(self):
        self.assertTrue(hasattr(QUERY, "run_query"))
        with self.assertRaisesRegex(ValueError, "not allowed"):
            QUERY.run_query(self.path, {"entity": "vm", "select": ["annotation"]})

    def test_reuses_sqlite_index_with_private_posix_permissions(self):
        self.assertTrue(hasattr(QUERY, "run_query"))
        index_path = Path(self.tempdir.name) / "inventory.sqlite"
        plan = {"entity": "vm", "metrics": ["count"]}

        first = QUERY.run_query(self.path, plan, index_path=index_path)
        second = QUERY.run_query(self.path, plan, index_path=index_path)

        self.assertFalse(first["index_reused"])
        self.assertTrue(second["index_reused"])
        self.assertTrue(index_path.is_file())
        # Windows uses inherited ACLs; st_mode cannot express owner-only access.
        if os.name != "nt":
            self.assertEqual(index_path.stat().st_mode & 0o777, 0o600)
        self.assertEqual(second["rows"], [{"count": 3}])

        workbook = load_workbook(self.path)
        workbook["vInfo"].append(
            ["new-vm", "poweredOn", False, 1, 1024, 1024, 512, "cluster-b", "esx-03", "vmx-19", False, "Ubuntu Linux", ""]
        )
        workbook.save(self.path)
        workbook.close()
        rebuilt = QUERY.run_query(self.path, plan, index_path=index_path)

        self.assertFalse(rebuilt["index_reused"])
        self.assertEqual(rebuilt["rows"], [{"count": 4}])

    def test_returns_only_warnings_relevant_to_the_query_entity(self):
        workbook = Workbook()
        workbook.remove(workbook.active)
        add_sheet(
            workbook,
            "vInfo",
            ["VM", "Powerstate", "Template", "CPUs", "Memory", "Cluster"],
            [["orphan-vm", "poweredOn", False, 2, 4096, "cluster-without-hosts"]],
        )
        add_sheet(workbook, "vHost", ["Host", "Cluster", "# Cores", "# Memory"], [])
        workbook.save(self.path)
        workbook.close()

        vm_result = QUERY.run_query(self.path, {"entity": "vm", "metrics": ["count"]})
        cluster_result = QUERY.run_query(
            self.path,
            {"entity": "cluster", "select": ["cluster", "cpu_ratio", "memory_ratio"]},
        )

        self.assertEqual(vm_result["warnings"], [])
        self.assertEqual(
            {item["code"] for item in cluster_result["warnings"]},
            {"missing_physical_cpu_capacity", "missing_physical_memory_capacity"},
        )

    def test_scopes_cluster_capacity_warnings_to_an_exact_cluster_filter(self):
        workbook = Workbook()
        workbook.remove(workbook.active)
        add_sheet(
            workbook,
            "vInfo",
            ["VM", "Powerstate", "Template", "CPUs", "Memory", "Cluster"],
            [
                ["valid-vm", "poweredOn", False, 2, 4096, "valid-cluster"],
                ["orphan-vm", "poweredOn", False, 2, 4096, "missing-cluster"],
            ],
        )
        add_sheet(
            workbook,
            "vHost",
            ["Host", "Cluster", "# Cores", "# Memory"],
            [["esx-01", "valid-cluster", 8, 32768]],
        )
        workbook.save(self.path)
        workbook.close()

        valid = QUERY.run_query(
            self.path,
            {"entity": "cluster", "metrics": ["count"], "filters": ["cluster=valid-cluster"]},
        )
        missing = QUERY.run_query(
            self.path,
            {"entity": "cluster", "metrics": ["count"], "filters": ["cluster=missing-cluster"]},
        )

        self.assertEqual(valid["warnings"], [])
        self.assertEqual(len(missing["warnings"]), 2)
        self.assertTrue(all(item["scope"] == "missing-cluster" for item in missing["warnings"]))

    def test_cli_emits_json_for_a_structured_query(self):
        completed = subprocess.run(
            [
                sys.executable,
                str(MODULE_PATH),
                str(self.path),
                "--entity",
                "vm",
                "--metric",
                "count",
                "--filter",
                "guest_os_family=linux",
                "--group-by",
                "cluster",
            ],
            check=False,
            capture_output=True,
            text=True,
        )

        self.assertEqual(completed.returncode, 0, completed.stderr)
        self.assertTrue(completed.stdout.strip())
        payload = json.loads(completed.stdout)
        self.assertEqual(payload["rows"], [{"cluster": "cluster-a", "count": 1}])

    def test_cli_rejects_a_non_allowlisted_field(self):
        completed = subprocess.run(
            [
                sys.executable,
                str(MODULE_PATH),
                str(self.path),
                "--entity",
                "vm",
                "--select",
                "annotation",
            ],
            check=False,
            capture_output=True,
            text=True,
        )

        self.assertNotEqual(completed.returncode, 0)
        self.assertIn("not allowed", completed.stderr)


if __name__ == "__main__":
    unittest.main()
