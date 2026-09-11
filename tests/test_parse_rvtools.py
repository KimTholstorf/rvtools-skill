import importlib.util
import json
import os
import subprocess
import sys
import tempfile
import unittest
from datetime import datetime, timedelta
from pathlib import Path
from unittest.mock import Mock, patch

from openpyxl import Workbook, load_workbook


MODULE_PATH = Path(__file__).parents[1] / "scripts" / "parse_rvtools.py"
SPEC = importlib.util.spec_from_file_location("parse_rvtools", MODULE_PATH)
PARSER = importlib.util.module_from_spec(SPEC)
assert SPEC.loader is not None
SPEC.loader.exec_module(PARSER)


def add_sheet(workbook, name, headers, rows):
    worksheet = workbook.create_sheet(name)
    worksheet.append(headers)
    for row in rows:
        worksheet.append(row)


class ParseInventoryTests(unittest.TestCase):
    def setUp(self):
        self.tempdir = tempfile.TemporaryDirectory()
        self.path = Path(self.tempdir.name) / "synthetic-rvtools.xlsx"

        workbook = Workbook()
        workbook.remove(workbook.active)
        add_sheet(
            workbook,
            "vInfo",
            [
                "VM",
                "Powerstate",
                "Template",
                "CPUs",
                "Memory",
                "Provisioned MiB",
                "In Use MiB",
                "Consolidation Needed",
                "HW version",
                "FT State",
                "Annotation",
                "OS according to the configuration file",
                "Cluster",
                "Host",
            ],
            [
                ["app-01", "poweredOn", False, 4, 8192, 102400, 51200, False, "vmx-13", "notConfigured", "", "Ubuntu Linux", "cluster-a", "esx-01"],
                ["db-ora-01", "suspended", False, 160, 2097152, 409600, 307200, True, "vmx-10", "running", "password=secret", "Oracle Linux 8", "cluster-a", "esx-02"],
                ["template-01", "poweredOff", True, 2, 4096, 10240, 5120, False, "vmx-19", "notConfigured", "", "Windows Server", "cluster-a", "esx-01"],
            ],
        )
        add_sheet(
            workbook,
            "vHost",
            [
                "Host",
                "Cluster",
                "CPU Model",
                "ESX Version",
                "CPU usage %",
                "Memory usage %",
                "# Cores",
                "# Memory",
                "Vendor",
                "Model",
                "HT Available",
                "HT Active",
                "# CPU",
                "Cores per CPU",
                "Speed",
                "BIOS Vendor",
                "BIOS Version",
                "BIOS Date",
            ],
            [
                ["esx-01", "cluster-a", "Intel Xeon Gold", "8.0.3", 32, 61, 16, 65536, "Cisco Systems Inc", "HXAF240C-M5SX", True, True, 2, 8, 3000, "Cisco Systems, Inc.", "C240M5.4.1.3m", "2022-07-08"],
                ["esx-02", "cluster-a", "AMD EPYC", "7.0.3", 81, 92, 16, 65536, "Dell Inc.", "PowerEdge R7525", True, False, 2, 8, 2800, "Dell Inc.", "2.19.1", "2024-01-15"],
            ],
        )
        add_sheet(
            workbook,
            "vDatastore",
            ["Name", "Capacity MiB", "Provisioned MiB", "In Use MiB", "Free MiB", "Free %", "Accessible"],
            [["ds-01", 1000000, 1100000, 950000, 50000, 5, True]],
        )
        add_sheet(
            workbook,
            "vSnapshot",
            ["VM", "Name", "Description", "Date / time", "Size MiB (total)"],
            [
                ["app-01", "old snapshot", "token: abc", datetime.now() - timedelta(days=45), 20480],
                ["app-01", "four-day snapshot", "", datetime.now() - timedelta(days=4), 1024],
                ["app-01", "recent snapshot", "", datetime.now() - timedelta(days=2), 512],
            ],
        )
        add_sheet(
            workbook,
            "vDisk",
            ["VM", "Template", "Disk", "Capacity MiB", "Raw", "Disk Mode", "Sharing mode", "Raw Comp. Mode", "Shared Bus"],
            [
                ["db-ora-01", False, "Hard disk 1", 204800, True, "independent_persistent", "sharingMultiWriter", "physicalMode", "physicalSharing"],
                ["app-01", False, "Hard disk 1", 51200, False, "persistent", "sharingNone", "", "noSharing"],
                ["template-01", True, "Hard disk 1", 10240, True, "independent_persistent", "sharingMultiWriter", "physicalMode", "physicalSharing"],
            ],
        )
        add_sheet(
            workbook,
            "vNetwork",
            ["VM", "Template", "Network", "Switch", "Connected", "Host"],
            [
                ["app-01", False, "dv-app", "dvSwitch-01", True, "esx-01"],
                ["db-ora-01", False, "Management Network", "vSwitch0", True, "esx-02"],
                ["app-01", False, "backup-vmk", "vSwitch1", True, "esx-01"],
                ["template-01", True, "Management Network", "vSwitch2", True, "esx-02"],
            ],
        )
        add_sheet(
            workbook,
            "vSC_VMK",
            ["Host", "Port Group"],
            [["esx-02", "Management Network"], ["esx-02", "backup-vmk"]],
        )
        add_sheet(
            workbook,
            "vUSB",
            ["VM", "Template", "Connected", "Device Type"],
            [["app-01", False, True, "USB device"], ["template-01", True, True, "USB device"]],
        )
        add_sheet(
            workbook,
            "vCD",
            ["VM", "Template", "Connected", "Starts Connected"],
            [
                ["app-01", False, True, True],
                ["offline-01", False, False, True],
                ["template-01", True, True, True],
            ],
        )
        add_sheet(
            workbook,
            "vTools",
            ["VM", "Powerstate", "Template", "Tools", "Tools Version", "Required Version"],
            [
                ["app-01", "poweredOn", False, "toolsNotRunning", "123", "124"],
                ["offline-01", "poweredOff", False, "toolsNotRunning", "123", "124"],
                ["template-01", "poweredOff", True, "toolsNotRunning", "123", "124"],
            ],
        )
        add_sheet(
            workbook,
            "dvPort",
            ["Port", "Type", "VLAN", "Allow Promiscuous", "Mac Changes", "Forged Transmits", "Blocked"],
            [
                ["dv-app", "ephemeral", 0, True, True, True, False],
                ["dv-untagged", "earlyBinding", None, False, False, False, False],
            ],
        )
        add_sheet(workbook, "vCluster", ["Name", "NumHosts"], [["cluster-a", 2]])
        add_sheet(workbook, "vHealth", ["Name", "Message", "Message type"], [["storage", "Path degraded", "Error"]])
        add_sheet(workbook, "vMetaData", ["RVTools version", "xlsx creation datetime"], [["4.7.1", datetime.now()]])
        workbook.save(self.path)
        workbook.close()

    def tearDown(self):
        self.tempdir.cleanup()

    def test_extracts_normalized_inventory_and_capacity(self):
        result = PARSER.analyze_workbook(self.path)

        self.assertEqual(result["schema_version"], "3.0")
        self.assertEqual(result["source"]["file"], self.path.name)
        self.assertEqual(result["source"]["rvtools_version"], "4.7.1")
        self.assertEqual(result["inventory"]["vms"], 2)
        self.assertEqual(result["inventory"]["templates"], 1)
        self.assertEqual(result["inventory"]["hosts"], 2)
        self.assertEqual(result["inventory"]["snapshots"], 3)
        self.assertEqual(result["inventory"]["virtual_disks"], 2)
        self.assertEqual(result["inventory"]["template_virtual_disks"], 1)
        self.assertEqual(result["inventory"]["network_adapters"], 3)
        self.assertEqual(result["inventory"]["template_network_adapters"], 1)
        self.assertEqual(result["capacity_mib"]["vm_provisioned"], 512000)
        self.assertEqual(result["capacity_mib"]["datastore_free"], 50000)
        self.assertEqual(result["facts"]["esxi_versions"], {"7.0.3": 1, "8.0.3": 1})

    def test_analysis_closes_source_workbook(self):
        workbook = load_workbook(self.path, read_only=True, data_only=True)
        close_spy = Mock(wraps=workbook.close)
        workbook.close = close_spy

        try:
            with patch.object(PARSER, "load_workbook", return_value=workbook):
                PARSER.analyze_workbook(self.path)
            close_spy.assert_called_once_with()
        finally:
            if not close_spy.called:
                workbook.close()

    def test_preserves_duplicate_headers_and_prefers_the_last_populated_value(self):
        duplicate_path = Path(self.tempdir.name) / "duplicate-headers.xlsx"
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "vDisk"
        worksheet.append(["VM", "Cluster", "Datacenter", "Cluster", "Host"])
        worksheet.append(["vm-01", "custom-tag-value", "dc-01", "actual-cluster", "esx-01"])
        workbook.save(duplicate_path)
        workbook.close()

        workbook = load_workbook(duplicate_path, read_only=True, data_only=True)
        try:
            row = PARSER._read_rows(workbook, "vDisk")[0]
        finally:
            workbook.close()

        self.assertEqual(row["cluster"], "custom-tag-value")
        self.assertEqual(row["cluster2"], "actual-cluster")
        self.assertEqual(PARSER._value(row, "Cluster"), "actual-cluster")

    def test_extracts_host_hardware_lifecycle_and_hyperthreading_facts(self):
        result = PARSER.analyze_workbook(self.path)

        self.assertEqual(
            result["facts"]["host_hardware_models"],
            [
                {"vendor": "Cisco Systems Inc", "model": "HXAF240C-M5SX", "hosts": 1},
                {"vendor": "Dell Inc.", "model": "PowerEdge R7525", "hosts": 1},
            ],
        )
        self.assertEqual(
            result["facts"]["host_hyperthreading"],
            {
                "available": {"true": 2, "false": 0, "unknown": 0},
                "active": {"true": 1, "false": 1, "unknown": 0},
            },
        )

        detections = {item["id"]: item for item in result["detections"]}
        self.assertEqual(detections["host_hyperthreading_inactive"]["count"], 1)
        self.assertEqual(
            detections["host_hyperthreading_inactive"]["examples"],
            [
                {
                    "host": "esx-02",
                    "vendor": "Dell Inc.",
                    "model": "PowerEdge R7525",
                    "ht_available": True,
                    "ht_active": False,
                }
            ],
        )

    def test_emits_queryable_guest_os_and_cluster_dimensions(self):
        workbook = load_workbook(self.path)
        vinfo = workbook["vInfo"]
        vinfo["O1"] = "OS according to the VMware Tools"
        vinfo["O2"] = "Microsoft Windows Server 2022 (64-bit)"
        workbook.save(self.path)
        workbook.close()

        result = PARSER.analyze_workbook(self.path)

        self.assertIn("guest_os_families", result["facts"])
        self.assertEqual(result["facts"]["guest_os_families"], {"linux": 1, "windows": 1})
        self.assertEqual(
            result["facts"]["guest_os_sources"],
            {"configuration": 1, "vmware_tools": 1},
        )
        self.assertEqual(result["facts"]["vm_counts_by_cluster"], {"cluster-a": 2})

    def test_calculates_cpu_and_memory_overcommit_per_cluster_and_overall(self):
        workbook = load_workbook(self.path)
        vinfo = workbook["vInfo"]
        vinfo["B2"] = "poweredOn"
        vinfo["D2"] = 32
        vinfo["E2"] = 65536
        vinfo["B3"] = "poweredOff"
        vinfo["D3"] = 16
        vinfo["E3"] = 32768
        vinfo.append(
            [
                "app-b",
                "poweredOn",
                False,
                16,
                65536,
                102400,
                51200,
                False,
                "vmx-19",
                "notConfigured",
                "",
                "Ubuntu Linux",
                "cluster-b",
                "esx-03",
            ]
        )
        workbook["vHost"].append(
            ["esx-03", "cluster-b", "Intel Xeon Gold", "8.0.3", 20, 30, 8, 32768]
        )
        workbook.save(self.path)
        workbook.close()

        result = PARSER.analyze_workbook(self.path)

        self.assertIn("overcommit", result)
        overcommit = result["overcommit"]
        self.assertEqual(overcommit["methodology"]["vm_scope"], "powered_on_workloads")
        clusters = {item["cluster"]: item for item in overcommit["clusters"]}
        self.assertEqual(set(clusters), {"cluster-a", "cluster-b"})
        self.assertEqual(
            clusters["cluster-a"],
            {
                "cluster": "cluster-a",
                "host_count": 2,
                "powered_on_vms": 1,
                "configured_vcpus": 32,
                "physical_cores": 32,
                "cpu_ratio": 1.0,
                "configured_memory_mib": 65536,
                "physical_memory_mib": 131072,
                "memory_ratio": 0.5,
                "powered_off_vms": 1,
                "powered_off_vcpus": 16,
                "powered_off_memory_mib": 32768,
            },
        )
        self.assertEqual(clusters["cluster-b"]["cpu_ratio"], 2.0)
        self.assertEqual(clusters["cluster-b"]["memory_ratio"], 2.0)
        self.assertEqual(overcommit["overall"]["configured_vcpus"], 48)
        self.assertEqual(overcommit["overall"]["physical_cores"], 40)
        self.assertEqual(overcommit["overall"]["cpu_ratio"], 1.2)
        self.assertEqual(overcommit["overall"]["memory_ratio"], 0.8)
        self.assertEqual(overcommit["overall"]["powered_off_vms"], 1)

    def test_warns_when_overcommit_host_capacity_is_unavailable(self):
        workbook = load_workbook(self.path)
        vhost = workbook["vHost"]
        vhost["G2"] = 0
        vhost["G3"] = 0
        vhost["H2"] = None
        vhost["H3"] = None
        workbook.save(self.path)
        workbook.close()

        result = PARSER.analyze_workbook(self.path)

        cluster = result["overcommit"]["clusters"][0]
        self.assertIsNone(cluster["cpu_ratio"])
        self.assertIsNone(cluster["memory_ratio"])
        self.assertIn("warnings", result["overcommit"])
        warning_codes = {item["code"] for item in result["overcommit"]["warnings"]}
        self.assertEqual(
            warning_codes,
            {"missing_physical_cpu_capacity", "missing_physical_memory_capacity"},
        )

    def test_warns_when_overcommit_aggregates_multiple_vcenters(self):
        workbook = load_workbook(self.path)
        vhost = workbook["vHost"]
        vhost["I1"] = "VI SDK Server"
        vhost["I2"] = "vcenter-a.example"
        vhost["I3"] = "vcenter-b.example"
        workbook.save(self.path)
        workbook.close()

        result = PARSER.analyze_workbook(self.path)

        warning_codes = {item["code"] for item in result["overcommit"]["warnings"]}
        self.assertIn("multi_vcenter_sources", warning_codes)

    def test_suppresses_overall_ratio_when_a_powered_cluster_lacks_host_capacity(self):
        workbook = load_workbook(self.path)
        workbook["vInfo"].append(
            [
                "orphan-cluster-vm",
                "poweredOn",
                False,
                8,
                16384,
                102400,
                51200,
                False,
                "vmx-19",
                "notConfigured",
                "",
                "Ubuntu Linux",
                "cluster-without-hosts",
                "missing-host",
            ]
        )
        workbook.save(self.path)
        workbook.close()

        result = PARSER.analyze_workbook(self.path)

        self.assertIsNone(result["overcommit"]["overall"]["cpu_ratio"])
        self.assertIsNone(result["overcommit"]["overall"]["memory_ratio"])

    def test_rejects_non_rvtools_workbook(self):
        invalid_path = Path(self.tempdir.name) / "invalid.xlsx"
        workbook = Workbook()
        workbook.save(invalid_path)
        workbook.close()

        with self.assertRaisesRegex(ValueError, "vInfo"):
            PARSER.analyze_workbook(invalid_path)

    def test_warns_when_optional_analysis_sheets_are_missing(self):
        minimal_path = Path(self.tempdir.name) / "minimal.xlsx"
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "vInfo"
        worksheet.append(["VM", "Powerstate", "Template"])
        worksheet.append(["vm-01", "poweredOn", False])
        workbook.save(minimal_path)
        workbook.close()

        result = PARSER.analyze_workbook(minimal_path)

        missing = {warning["sheet"] for warning in result["warnings"] if warning["code"] == "missing_sheet"}
        self.assertIn("vHost", missing)
        self.assertIn("vDisk", missing)

    def test_detects_vm_configuration_and_guest_risks(self):
        result = PARSER.analyze_workbook(self.path)
        detections = {item["id"]: item for item in result["detections"]}

        expected = {
            "vm_suspended",
            "consolidation_needed",
            "vm_cpu_large",
            "vm_memory_large",
            "legacy_vm_hardware",
            "fault_tolerance_enabled",
            "oracle_workload",
            "possible_cleartext_secret",
            "tools_not_running",
            "connected_usb",
            "connected_cdrom",
            "cdrom_starts_connected",
        }
        self.assertTrue(expected.issubset(detections))
        self.assertEqual(detections["vm_suspended"]["severity"], "high")
        self.assertEqual(detections["vm_suspended"]["count"], 1)
        self.assertEqual(detections["legacy_vm_hardware"]["count"], 2)
        self.assertEqual(detections["tools_not_running"]["count"], 1)
        self.assertEqual(detections["connected_usb"]["count"], 1)
        self.assertEqual(detections["connected_cdrom"]["count"], 1)
        self.assertEqual(detections["connected_cdrom"]["severity"], "medium")
        self.assertEqual(detections["cdrom_starts_connected"]["count"], 1)
        self.assertNotIn("lenses", detections["vm_suspended"])
        self.assertIn("source.power_state", detections["vm_suspended"]["tags"])
        self.assertNotIn("password=", str(detections["possible_cleartext_secret"]["examples"]).casefold())

    def test_detects_vms_provisioned_above_ten_tib(self):
        workbook = load_workbook(self.path)
        worksheet = workbook["vInfo"]
        worksheet["F2"] = 10 * 1024 * 1024
        worksheet["F3"] = 10 * 1024 * 1024 + 1
        workbook.save(self.path)
        workbook.close()

        result = PARSER.analyze_workbook(self.path)
        detections = {item["id"]: item for item in result["detections"]}

        self.assertIn("vm_provisioned_storage_large", detections)
        self.assertEqual(detections["vm_provisioned_storage_large"]["count"], 1)
        self.assertEqual(detections["vm_provisioned_storage_large"]["examples"][0]["vm"], "db-ora-01")

    def test_detects_storage_and_snapshot_risks(self):
        result = PARSER.analyze_workbook(self.path)
        detections = {item["id"]: item for item in result["detections"]}

        expected = {
            "stale_snapshot",
            "raw_device_mapping",
            "independent_disk",
            "shared_disk",
            "datastore_low_free_space",
            "datastore_overprovisioned",
        }
        self.assertTrue(expected.issubset(detections))
        self.assertGreaterEqual(detections["stale_snapshot"]["examples"][0]["age_days"], 44)
        self.assertEqual(detections["stale_snapshot"]["count"], 2)
        self.assertNotIn("snapshot", detections["stale_snapshot"]["examples"][0])
        self.assertIn("created_at", detections["stale_snapshot"]["examples"][0])
        self.assertEqual(detections["raw_device_mapping"]["examples"][0]["compatibility_mode"], "physicalMode")
        self.assertEqual(detections["raw_device_mapping"]["count"], 1)
        self.assertEqual(detections["independent_disk"]["count"], 1)
        self.assertEqual(detections["shared_disk"]["count"], 1)

    def test_detects_network_migration_risks(self):
        result = PARSER.analyze_workbook(self.path)
        detections = {item["id"]: item for item in result["detections"]}

        expected = {
            "standard_vswitch_attachment",
            "vmkernel_network_attachment",
            "dvport_vlan_zero",
            "dvport_promiscuous_mode",
            "dvport_mac_changes",
            "dvport_forged_transmits",
            "dvport_ephemeral_binding",
        }
        self.assertTrue(expected.issubset(detections))
        self.assertEqual(detections["vmkernel_network_attachment"]["examples"][0]["network"], "Management Network")
        self.assertEqual(detections["vmkernel_network_attachment"]["count"], 1)
        self.assertEqual(detections["standard_vswitch_attachment"]["count"], 2)
        self.assertEqual(detections["dvport_vlan_zero"]["count"], 2)
        self.assertEqual(detections["dvport_mac_changes"]["count"], 1)

    def test_detects_host_and_environment_health_risks(self):
        result = PARSER.analyze_workbook(self.path)
        detections = {item["id"]: item for item in result["detections"]}

        expected = {
            "non_intel_host",
            "esxi_version_spread",
            "host_cpu_pressure",
            "host_memory_pressure",
            "rvtools_health_error",
        }
        self.assertTrue(expected.issubset(detections))
        self.assertEqual(detections["esxi_version_spread"]["count"], 2)

    def test_cli_emits_json_and_bounds_examples(self):
        completed = subprocess.run(
            [sys.executable, str(MODULE_PATH), str(self.path), "--max-examples", "1"],
            check=False,
            capture_output=True,
            text=True,
        )

        self.assertEqual(completed.returncode, 0, completed.stderr)
        self.assertTrue(completed.stdout.strip())
        payload = json.loads(completed.stdout)
        self.assertEqual(payload["schema_version"], "3.0")
        self.assertTrue(all(len(item["examples"]) <= 1 for item in payload["detections"]))

    def test_adds_a_selected_target_migration_assessment(self):
        result = PARSER.analyze_workbook(self.path, target="avs")

        self.assertEqual(result["migration"]["target"]["id"], "avs")
        self.assertEqual(result["migration"]["scope"]["workload_vms"], 2)
        self.assertEqual(result["sizing"]["status"], "attention_required")
        self.assertEqual(result["sizing"]["policy"]["id"], "recommended")
        self.assertEqual(result["sizing"]["topology"]["id"], "consolidated")

    def test_accepts_an_explicit_active_only_source_aligned_sizing_policy(self):
        result = PARSER.analyze_workbook(
            self.path,
            target="ocvs",
            sizing_policy="active_only",
            sizing_topology="source_aligned",
            primary_source_cluster="cluster-a",
        )

        self.assertEqual(result["sizing"]["status"], "complete")
        self.assertEqual(result["sizing"]["policy"]["id"], "active_only")
        self.assertEqual(result["sizing"]["topology"]["id"], "source_aligned")
        self.assertEqual(result["sizing"]["topology"]["primary_target_cluster"], "cluster-a")
        by_vm = {row["vm"]: row for row in result["migration"]["vm_methods"]}
        self.assertEqual(by_vm["db-ora-01"]["methods"]["rav"]["status"], "blocked")

    def test_multi_cluster_target_requires_a_topology_choice(self):
        workbook = load_workbook(self.path)
        workbook["vInfo"].append(
            [
                "app-02",
                "poweredOn",
                False,
                2,
                4096,
                10240,
                5120,
                False,
                "vmx-19",
                "notConfigured",
                "",
                "Ubuntu Linux",
                "cluster-b",
                "esx-03",
            ]
        )
        workbook["vHost"].append(
            [
                "esx-03",
                "cluster-b",
                "Intel Xeon Gold",
                "8.0.3",
                20,
                30,
                16,
                65536,
                "Dell Inc.",
                "PowerEdge R760",
                True,
                True,
                2,
                8,
                2800,
                "Dell Inc.",
                "1.0",
                "2026-01-01",
            ]
        )
        workbook.save(self.path)
        workbook.close()

        result = PARSER.analyze_workbook(self.path, target="ocvs")

        self.assertEqual(result["sizing"]["status"], "topology_selection_required")
        self.assertEqual(result["sizing"]["default_policy_id"], "recommended")
        self.assertEqual(
            result["sizing"]["available_topologies"],
            ["consolidated", "source_aligned"],
        )

    def test_cli_accepts_a_migration_target(self):
        completed = subprocess.run(
            [sys.executable, str(MODULE_PATH), str(self.path), "--target", "gcve"],
            check=False,
            capture_output=True,
            text=True,
        )

        self.assertEqual(completed.returncode, 0, completed.stderr)
        payload = json.loads(completed.stdout)
        self.assertEqual(payload["migration"]["target"]["id"], "gcve")

    def test_cli_refuses_to_overwrite_source_workbook(self):
        original = self.path.read_bytes()

        completed = subprocess.run(
            [sys.executable, str(MODULE_PATH), str(self.path), "--output", str(self.path)],
            check=False,
            capture_output=True,
            text=True,
        )

        self.assertNotEqual(completed.returncode, 0)
        self.assertIn("source workbook", completed.stderr)
        self.assertEqual(self.path.read_bytes(), original)

    def test_cli_writes_output_and_uses_private_posix_permissions(self):
        output_path = Path(self.tempdir.name) / "analysis.json"

        completed = subprocess.run(
            [sys.executable, str(MODULE_PATH), str(self.path), "--output", str(output_path)],
            check=False,
            capture_output=True,
            text=True,
        )

        self.assertEqual(completed.returncode, 0, completed.stderr)
        self.assertTrue(output_path.is_file())
        # Windows uses inherited ACLs; st_mode cannot express owner-only access.
        if os.name != "nt":
            self.assertEqual(output_path.stat().st_mode & 0o777, 0o600)

    def test_cli_writes_output_when_fchmod_is_unavailable(self):
        output_path = Path(self.tempdir.name) / "portable-analysis.json"
        had_fchmod = hasattr(PARSER.os, "fchmod")
        original_fchmod = getattr(PARSER.os, "fchmod", None)
        if had_fchmod:
            delattr(PARSER.os, "fchmod")

        try:
            self.assertEqual(
                PARSER.main([str(self.path), "--output", str(output_path)]),
                0,
            )
        finally:
            if had_fchmod:
                PARSER.os.fchmod = original_fchmod

        self.assertTrue(output_path.is_file())


if __name__ == "__main__":
    unittest.main()
